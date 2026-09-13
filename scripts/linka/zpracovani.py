"""Příprava úlohy: stažení, kontrola struktury a zpracování do pracovního adresáře.

Nic se nezapisuje do public/. Výstupy leží v data/linka/prace/<KÓD>/ a do repozitáře
se dostanou až předáním po schválení.
"""
from __future__ import annotations

import json
import os
import statistics
import subprocess
import sys
from pathlib import Path

from . import jadro

# Sloupce, bez kterých zpracovatelé dat o uchazečích nefungují.
POVINNE_UCHAZECI = (
    [f"ss{k}_redizo" for k in range(1, 6)] + [f"ss{k}_kkov" for k in range(1, 6)]
    + [f"ss{k}_prijat" for k in range(1, 6)] + [f"ss{k}_duvod_neprijeti" for k in range(1, 6)]
    + ["c_m_procentni_skor"]
)


def prozkoumej_xlsx(soubor: Path, pocitat_radky: bool = True) -> dict:
    """Listy, hlavička prvního listu a počet řádků."""
    import openpyxl

    wb = openpyxl.load_workbook(soubor, read_only=True)
    ws = wb.worksheets[0]
    hlavicka: list[str] = []
    radku = 0
    for i, radek in enumerate(ws.iter_rows(values_only=True)):
        if not hlavicka:
            hodnoty = [str(v).strip() for v in radek if v not in (None, "")]
            if len(hodnoty) >= 3:
                hlavicka = [str(v).strip() if v is not None else "" for v in radek]
            elif i > 5:
                break
            continue
        if not pocitat_radky:
            break
        radku += 1
    return {"listy": wb.sheetnames, "hlavicka": hlavicka, "sloupcu": len(hlavicka), "radku": radku}


def rozdil_struktury(novy: dict, stary: dict | None) -> list[str]:
    if not stary:
        return []
    zmeny = []
    if novy["listy"] != stary["listy"]:
        zmeny.append(f"listy se změnily z {stary['listy']} na {novy['listy']}")
    pridane = [s for s in novy["hlavicka"] if s and s not in stary["hlavicka"]]
    odebrane = [s for s in stary["hlavicka"] if s and s not in novy["hlavicka"]]
    if not pridane and not odebrane and [x for x in novy["hlavicka"] if x] != [x for x in stary["hlavicka"] if x]:
        zmeny.append("pořadí sloupců se změnilo; skripty čtoucí sloupce podle pozice dají chybné výsledky")
    if pridane:
        zmeny.append(f"přibyly sloupce {pridane[:8]}")
    if odebrane:
        zmeny.append(f"zmizely sloupce {odebrane[:8]}")
    return zmeny


def predchozi_soubor(sada: dict) -> Path | None:
    soubor = sada.get("zobrazeno", {}).get("soubor")
    if soubor and not soubor.startswith("http"):
        p = jadro.KOREN / soubor
        if p.is_file() and p.suffix == ".xlsx":
            return p
    return None


# ---------------------------------------------------------------- zpracovatelé

def spust(skript: str, *argumenty: str) -> str:
    vysledek = subprocess.run(
        [sys.executable, str(jadro.KOREN / "scripts" / skript), *argumenty],
        capture_output=True, text=True, cwd=jadro.KOREN,
    )
    if vysledek.returncode != 0:
        raise RuntimeError(f"{skript} skončil chybou: {vysledek.stderr.strip()[-400:]}")
    return vysledek.stdout.strip()


def srovnej_pasma(novy: Path, stavajici: Path) -> dict:
    n = json.loads(novy.read_text(encoding="utf-8"))["data"]
    s = json.loads(stavajici.read_text(encoding="utf-8"))["data"] if stavajici.exists() else {}
    spolecne = [k for k in n if k in s]
    zmena = [abs(n[k]["min_prijaty"] - s[k]["min_prijaty"]) for k in spolecne]
    return {
        "oboru_nove": len(n), "oboru_na_webu": len(s),
        "s_pasmy_nove": sum(1 for v in n.values() if "pasma" in v),
        "s_pasmy_na_webu": sum(1 for v in s.values() if "pasma" in v),
        "spolecnych_oboru": len(spolecne),
        "median_zmeny_nejnizsiho_prijateho": round(statistics.median(zmena), 1) if zmena else None,
    }


def zpracuj_uchazeci(uloha: dict, soubor: Path, prace: Path, struktura: dict) -> dict:
    chybi = [s for s in POVINNE_UCHAZECI if s not in struktura["hlavicka"]]
    if chybi:
        raise ValueError(f"chybí povinné sloupce {chybi}")
    rok = int(uloha["obdobi"])
    pasma = prace / f"pasma_prijeti_{rok}.json"
    soubeh = prace / f"soubeh_prihlasek_{rok}.json"
    vystup_pasma = spust("build-pasma-prijeti.py", "--zdroj", str(soubor), "--vystup", str(pasma), "--rok", str(rok))
    vystup_soubeh = spust("build-soubeh-prihlasek.py", "--zdroj", str(soubor), "--vystup", str(soubeh), "--rok", str(rok))
    srovnani = srovnej_pasma(pasma, jadro.KOREN / "public" / f"pasma_prijeti_{uloha['zobrazene_obdobi']}.json")
    # Pojistka proti tichému selhání: v září 2026 změnil CERMAT příznak přijetí z čísla na text
    # a zpracování vrátilo nula oborů, aniž by skončilo chybou.
    if srovnani["oboru_nove"] == 0:
        raise ValueError("zpracování nevrátilo žádný obor; zkontroluj formát příznaku přijetí a skóre")
    min_podil = float(os.environ.get("LINKA_MIN_PODIL_OBORU", "0.5"))
    if srovnani["oboru_na_webu"] and srovnani["oboru_nove"] < min_podil * srovnani["oboru_na_webu"]:
        raise ValueError(
            f"podezřele málo oborů: {srovnani['oboru_nove']} proti {srovnani['oboru_na_webu']} na webu"
        )
    return {
        "zpracovatel": "cermat-uchazeci-kolo1",
        "vystupy_skriptu": [vystup_pasma, vystup_soubeh],
        "srovnani": srovnani,
        "predani": {
            str(pasma): f"public/pasma_prijeti_{rok}.json",
            str(soubeh): f"public/soubeh_prihlasek_{rok}.json",
        },
        "dopad": (
            "Nové soubory za další rok; web je nečte, dokud se nepřepne registr a nenapojí kód."
            if uloha["druh"] == "nove_obdobi"
            else "Přepíše soubory, které web čte; po sloučení se změní čísla na stránkách oborů."
        ),
    }


ZPRACOVATELE = {"cermat-uchazeci-kolo1": zpracuj_uchazeci}


def priprav(uloha: dict, registr: dict, stahni_fn=jadro.stahni) -> None:
    """Připraví jednu úlohu a změní její stav na pripraveno, selhalo nebo bez_zmeny."""
    sada = registr["sady"][uloha["sada"]]
    prace = jadro.prace_cesta() / uloha["kod"]
    priprava: dict = {"cas": jadro.ted()}
    uloha["priprava"] = priprava

    if uloha["druh"] == "zmizelo":
        priprava["souhrn"] = f"Zdroj {uloha['url']} vrací 404."
        priprava["dopad"] = "Sada se nedá obnovit, dokud se nenajde nový zdroj."
        jadro.zmen_stav(uloha, "pripraveno")
        return

    try:
        nazev = uloha["url"].rsplit("/", 1)[-1] or "soubor"
        stazeno = stahni_fn(uloha["url"], prace / nazev)
        priprava["stazeno"] = {k: v for k, v in stazeno.items() if k != "soubor"}
        soubor = Path(stazeno["soubor"])

        # Revize, která obsahově nic nemění, se neoznamuje.
        znamy_otisk = sada.get("zobrazeno", {}).get("sha256")
        if znamy_otisk and znamy_otisk == stazeno["sha256"]:
            jadro.zmen_stav(uloha, "bez_zmeny", "otisk shodný se zobrazeným souborem")
            return

        if soubor.suffix == ".xlsx":
            struktura = prozkoumej_xlsx(soubor)
            predchozi = predchozi_soubor(sada)
            stara = prozkoumej_xlsx(predchozi, pocitat_radky=False) if predchozi else None
            priprava["struktura"] = {k: struktura[k] for k in ("listy", "sloupcu", "radku")}
            priprava["zmeny_struktury"] = rozdil_struktury(struktura, stara)
        else:
            struktura = None
            priprava["zmeny_struktury"] = []

        zpracovatel = ZPRACOVATELE.get(uloha["sada"])
        if zpracovatel and struktura is not None:
            priprava["zpracovani"] = zpracovatel(uloha, soubor, prace, struktura)
            priprava["dopad"] = priprava["zpracovani"]["dopad"]
        else:
            priprava["dopad"] = "Soubor je stažený a zkontrolovaný; sada nemá automatické zpracování. " + sada.get("aktualizace", {}).get("lidsky_krok", "")
        jadro.zmen_stav(uloha, "pripraveno")
    except Exception as e:  # jakákoli chyba přípravy se oznámí, linka nespadne
        priprava["chyba"] = str(e)
        jadro.zmen_stav(uloha, "selhalo", str(e)[:200])
