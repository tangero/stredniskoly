#!/usr/bin/env python3
"""Analytický prototyp maturitních výsledků (docs/maturitni-vysledky-a-kvalita-skoly-2027.md, §9.1).

Interní nástroj: nezapisuje do public/, nemění runtime ani veřejné rozhraní.

    python3 scripts/maturita_prototype.py --self-test
    python3 scripts/maturita_prototype.py --input-dir /cesta/k/xlsx \
        --output docs/podklady/maturita-prototyp-2026.json

Vstupní adresář má obsahovat soubory CERMAT:
    MZ{rok}j_SC_skolobory.xlsx            jaro 2015–2026 (společná část, školy a SMO16)
    MZ{rok}jap_SC_skolobory.xlsx          stav po podzimním období 2015–2025 (volitelné)
    JPZ{rok}_skoly-skolobory_vysledky.xlsx školní agregáty JPZ 2017–2023

Co skript dělá:
    1. manifest zdrojů (název, SHA-256, počty řádků podle úrovně agregace);
    2. dlouhou tabulku výsledků podle datového kontraktu (§4);
    3. kontextové srovnání uvnitř SMO16 s kvalifikátorem ze standardní chyby (§5.2);
    4. stabilitu v čase (§5.2);
    5. backtest vrstvy 3: model MZ ~ vstup JPZ + SMO16 + kraj + velikost, korelace reziduí
       školy mezi ročníky (§5.3).

Omezení: loader MZ hledá sloupce podle názvů v obou hlavičkových řádcích. Mapa sloupců
vznikla bez přístupu ke skutečnému souboru; při nenalezení povinných polí skript skončí
chybou a vypíše nalezenou hlavičku. Výsledek backtestu není důkaz, je to vstup pro
metodickou oponenturu.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import random
import re
import statistics
import subprocess
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE_BASE_MZ = 'https://data.cermat.cz/files/files/MZ/agregovana_data_skoly/'
SOURCE_BASE_JPZ = 'https://data.cermat.cz/files/files/JPZ/agregovana_data_skoly/'
MZ_YEARS = range(2015, 2027)
JPZ_YEARS = range(2017, 2024)
COHORT_SHIFT = 4          # čtyřleté obory: JPZ rok t → MZ rok t+4
MIN_N_REFERENCE = 10      # škola vstupuje do reference skupiny až od tohoto n
MIN_N_BACKTEST = 10
Z_95 = 1.96
PUBLISH_THRESHOLD_R = 0.3  # §5.3: kritérium zapsané předem
COVID_MZ_YEARS = {2021}
COVID_JPZ_YEARS = {2020, 2021}
SUBJECTS = ('spolecna_cast', 'cj', 'ma', 'aj', 'nj', 'fj', 'sj', 'ru')


# ---------------------------------------------------------------------------
# Pomocné funkce
# ---------------------------------------------------------------------------

def strip_accents(text: str) -> str:
    text = unicodedata.normalize('NFKD', str(text))
    return ''.join(c for c in text if not unicodedata.combining(c)).upper()


def sha256_of(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def as_number(value):
    """Nula, chybějící údaj a nezveřejněná hodnota jsou tři stavy; text nikdy nepřevádíme na 0."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    text = str(value).strip().replace(',', '.').replace('%', '')
    if not text or text in {'-', '–', 'x', 'X', '.', 'n/a', 'N/A'}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def as_count(value):
    number = as_number(value)
    if number is None:
        return None
    if number < 0 or int(number) != number:
        raise ValueError(f'Neplatný počet: {value!r}')
    return int(number)


def rate(numerator, denominator):
    if numerator is None or denominator is None or denominator <= 0:
        return None
    return round(numerator / denominator, 4)


def git_commit() -> str | None:
    try:
        return subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip()
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Loader MZ: objevování sloupců podle názvů
# ---------------------------------------------------------------------------

SUBJECT_PATTERNS = [
    ('cj', r'CESK|CJL|\bCJ\b'),
    ('ma', r'MATEMAT|\bMA\b'),
    ('aj', r'ANGLIC|\bAJ\b'),
    ('nj', r'NEMEC|\bNJ\b'),
    ('fj', r'FRANCOUZ|\bFJ\b'),
    ('sj', r'SPANEL|\bSJ\b'),
    ('ru', r'RUSK|\bRJ\b|\bRU\b'),
    ('spolecna_cast', r'SPOLECN|CELKEM|MATURITNI ZKOUSK|\bMZ\b'),
]

# Pořadí je důležité: specifičtější vzory dřív.
METRIC_PATTERNS = [
    ('standardDeviation', r'SMERODAT'),
    ('averagePercentile', r'PERCENTIL'),
    ('averagePercentScore', r'SKOR'),
    ('subjectChoiceShare', r'VOLB'),
    ('grossFailureRate', r'HRUB.*NEUSP'),
    ('netFailureRate', r'CIST.*NEUSP'),
    ('nonParticipationRate', r'(NEUCAST|NEKONAL|OMLUV).*(%|PODIL)|(%|PODIL).*(NEUCAST|NEKONAL|OMLUV)'),
    ('failRate', r'NEUSP.*(%|PODIL)|(%|PODIL).*NEUSP'),
    ('passRate', r'USPESNOST|USP.*(%|PODIL)|(%|PODIL).*USP'),
    ('absent', r'NEKONAL|NEUCAST|OMLUV|NEDOSTAV'),
    ('failed', r'NEUSP'),
    ('passed', r'USPEL|USPESN'),
    ('took', r'KONAL'),
    ('registered', r'PRIHLAS'),
]

REQUIRED_MZ = {('cj', 'averagePercentile')}
ID_PATTERNS = {
    'level': r'UROVEN|AGREGAC|TYP RADK|GRANULAR',
    'redizo': r'REDIZO|RED_IZO|RED IZO',
    'smo16': r'SMO16|SMO 16|SKUPINA OBOR',
    'smo16Label': r'NAZEV SKUP|SKUPINA OBORU - NAZ|SMO16 - NAZ|SMO16_NAZ',
    'kraj': r'\bKRAJ\b',
    'schoolType': r'TYP SKOLY|TYP_SKOLY',
    'name': r'NAZEV SKOLY|NAZEV_SKOLY|SKOLA',
}


def read_rows(path: Path, sheet: str | None = None) -> list[list]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = book[sheet] if sheet else book[book.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    book.close()
    return rows


def merged_headers(rows: list[list]) -> list[tuple[str, str]]:
    """Dvouřádková hlavička: první řádek (skupina/předmět) se doplní dopředu."""
    width = max(len(r) for r in rows[:2])
    top = list(rows[0]) + [None] * (width - len(rows[0]))
    bottom = list(rows[1]) + [None] * (width - len(rows[1]))
    filled, last = [], ''
    for value in top:
        if value not in (None, ''):
            last = str(value)
        filled.append(last)
    return [(filled[i], '' if bottom[i] is None else str(bottom[i])) for i in range(width)]


def classify_mz_columns(headers: list[tuple[str, str]]) -> dict:
    """Vrátí mapu {'ids': {role: index}, 'metrics': {(subject, metric): index}, 'unmatched': [...]}."""
    ids, metrics, unmatched = {}, {}, []
    for index, (group, label) in enumerate(headers):
        g, l = strip_accents(group), strip_accents(label)
        full = f'{g} | {l}'
        role = next((r for r, pat in ID_PATTERNS.items() if re.search(pat, l) or (not l and re.search(pat, g))), None)
        if role and role not in ids and not re.search(r'PRIHLAS|KONAL|USP|SKOR|PERCENTIL|SMERODAT|VOLB', full):
            ids[role] = index
            continue
        subject = next((s for s, pat in SUBJECT_PATTERNS if re.search(pat, g)), None)
        if subject is None:
            subject = next((s for s, pat in SUBJECT_PATTERNS if re.search(pat, l)), None)
        metric = next((m for m, pat in METRIC_PATTERNS if re.search(pat, l)), None)
        if subject and metric and (subject, metric) not in metrics:
            metrics[(subject, metric)] = index
        else:
            unmatched.append({'index': index, 'group': group, 'label': label})
    return {'ids': ids, 'metrics': metrics, 'unmatched': unmatched}


def detect_level(value) -> str | None:
    text = strip_accents(value or '').replace('-', '_').replace(' ', '_')
    if not text:
        return None
    if 'REDIZO_SMO16' in text or text in {'SKOLA_SMO16', 'SKOLA_SKUPINA'}:
        return 'redizo_smo16'
    if text in {'REDIZO', 'SKOLA'}:
        return 'redizo'
    return text.lower()


def load_mz(path: Path, year: int, period: str, source_url: str, lenient: bool = False) -> tuple[list[dict], dict]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((s for s in book.sheetnames if s.strip() == str(year)), book.sheetnames[0])
    population = book['vysvetlivky']['B2'].value if 'vysvetlivky' in book.sheetnames else None
    book.close()
    rows = read_rows(path, sheet)
    if len(rows) < 3:
        raise ValueError(f'{path.name}: příliš málo řádků')
    headers = merged_headers(rows)
    cmap = classify_mz_columns(headers)
    missing = [k for k in REQUIRED_MZ if k not in cmap['metrics']]
    if missing and not lenient:
        raise ValueError(
            f'{path.name}: nenalezeny povinné sloupce {missing}. Nalezená hlavička:\n'
            + '\n'.join(f'  [{i}] {g!r} | {l!r}' for i, (g, l) in enumerate(headers))
        )
    ids = cmap['ids']
    level_index = ids.get('level', 2)   # audit R1: úroveň agregace ve třetím sloupci
    redizo_index = ids.get('redizo', 0)
    sha = sha256_of(path)
    retrieved = datetime.now(timezone.utc).isoformat()
    records, level_counts = [], defaultdict(int)
    for row in rows[2:]:
        if not row or all(v in (None, '') for v in row):
            continue
        level = detect_level(row[level_index] if level_index < len(row) else None)
        level_counts[level or 'unknown'] += 1
        if level not in ('redizo', 'redizo_smo16'):
            continue
        redizo = str(row[redizo_index]).strip() if redizo_index < len(row) and row[redizo_index] is not None else ''
        redizo = re.sub(r'\.0$', '', redizo)
        if not re.fullmatch(r'\d{9}', redizo):
            continue
        smo16_raw = row[ids['smo16']] if 'smo16' in ids and ids['smo16'] < len(row) else None
        smo16 = int(as_number(smo16_raw)) if level == 'redizo_smo16' and as_number(smo16_raw) is not None else None
        smo16_label = str(row[ids['smo16Label']]) if 'smo16Label' in ids and row[ids['smo16Label']] is not None else (
            str(smo16_raw) if level == 'redizo_smo16' and smo16 is None and smo16_raw is not None else None)
        kraj = str(row[ids['kraj']]) if 'kraj' in ids and row[ids['kraj']] is not None else None
        school_type = str(row[ids['schoolType']]) if 'schoolType' in ids and row[ids['schoolType']] is not None else None
        for subject in SUBJECTS:
            fields = {m: cmap['metrics'].get((subject, m)) for m in dict(METRIC_PATTERNS)}
            if all(i is None for i in fields.values()):
                continue

            def get(metric, conv=as_number):
                index = fields.get(metric)
                return conv(row[index]) if index is not None and index < len(row) else None

            took = get('took', as_count)
            registered = get('registered', as_count)
            passed = get('passed', as_count)
            failed = get('failed', as_count)
            absent = get('absent', as_count)
            if took is None and registered is not None and absent is not None:
                took = registered - absent
            reasons = []
            if took is None:
                reasons.append('chybí počet konajících')
            elif took < MIN_N_REFERENCE:
                reasons.append(f'malý vzorek n={took}')
            if period == 'stav_po_podzimu':
                reasons.append('populace po podzimním období, nesčítat s jarem')
            if year in COVID_MZ_YEARS:
                reasons.append('mimořádný ročník 2021')
            if year <= 2020 and subject == 'spolecna_cast':
                reasons.append('společná část do 2020 zahrnovala písemnou práci a ústní zkoušku; jiná škála')
            percentile = get('averagePercentile')
            records.append({
                'id': f'{year}_{period}_{level}_{redizo}_{smo16 if smo16 is not None else "x"}_{subject}',
                'schoolYear': year, 'period': period, 'grain': level, 'redizo': redizo,
                'smo16': smo16, 'smo16Label': smo16_label, 'kraj': kraj, 'schoolType': school_type,
                'subject': subject,
                'registered': registered, 'took': took, 'passed': passed, 'failed': failed, 'absent': absent,
                'passRate': get('passRate') if get('passRate') is not None else rate(passed, took),
                'grossFailureRate': get('grossFailureRate'),
                'netFailureRate': get('netFailureRate') if get('netFailureRate') is not None else (
                    get('failRate') if get('failRate') is not None else rate(failed, took)),
                'nonParticipationRate': get('nonParticipationRate') if get('nonParticipationRate') is not None else rate(absent, registered),
                'averagePercentScore': get('averagePercentScore'),
                'standardDeviation': get('standardDeviation'),
                'averagePercentile': percentile,
                'subjectChoiceShare': get('subjectChoiceShare'),
                'populationNote': population or 'nezjištěno (list vysvětlivky chybí)',
                'sourceUrl': source_url, 'sourceFileSha256': sha, 'sourceValidAt': None, 'retrievedAt': retrieved,
                'quality': 'unavailable' if percentile is None and took is None else ('limited' if reasons else 'complete'),
                'qualityReasons': reasons,
            })
    manifest = {
        'file': path.name, 'url': source_url, 'sha256': sha, 'bytes': path.stat().st_size, 'sheet': sheet,
        'physical_rows': len(rows), 'aggregation_rows': dict(level_counts), 'population': population,
        'column_map': {f'{s}.{m}': i for (s, m), i in cmap['metrics'].items()},
        'id_columns': ids, 'unmatched_columns': cmap['unmatched'],
        'headers': [f'{g} | {l}' for g, l in headers],
    }
    return records, manifest


# ---------------------------------------------------------------------------
# Loader JPZ 2017–2023: školní agregáty konajících (ne přijatých)
# ---------------------------------------------------------------------------

def load_jpz(path: Path, year: int, source_url: str) -> tuple[list[dict], dict]:
    rows = read_rows(path)
    headers = merged_headers(rows)
    labels = [strip_accents(l) for _, l in headers]
    groups = [strip_accents(g) for g, _ in headers]

    def find(pattern, start=0):
        return next((i for i in range(start, len(labels)) if re.search(pattern, labels[i])), None)

    redizo_i = find(r'REDIZO') or 0
    group_i = find(r'OBOROV')
    rocnik_i = find(r'ROCNIK|DELKA')
    kraj_i = find(r'KRAJ.*KOD|KRAJ \(KOD\)')
    # Dva bloky (ČJ, MA): PŘIHLÁŠENI, KONALI, ..., PRŮMĚRNÉ PERCENTILOVÉ UMÍSTĚNÍ, SMĚRODATNÁ ODCHYLKA.
    # Pozor: 'SMĚRODATNÁ ODCHYLKA (PERCENTIL. UMÍSTĚNÍ)' také obsahuje slovo PERCENTIL.
    def find_percentile(start):
        return next((i for i in range(start, len(labels))
                     if re.search(r'PERCENTIL', labels[i]) and not re.search(r'SMERODAT|ODCHYLK', labels[i])), None)

    blocks = []
    start = 0
    while True:
        p = find_percentile(start)
        if p is None:
            break
        konali = next((i for i in range(p, -1, -1) if re.fullmatch(r'KONALI', labels[i])), None)
        subject = next((s for s, pat in SUBJECT_PATTERNS[:2] if re.search(pat, groups[p])), None)
        blocks.append({'subject': subject, 'percentile': p, 'took': konali,
                       'sd': find(r'SMERODAT', p)})
        start = p + 1
    if len(blocks) == 2 and any(b['subject'] is None for b in blocks):
        blocks[0]['subject'], blocks[1]['subject'] = 'cj', 'ma'   # předpoklad pořadí ČJ, MA; zapsáno do manifestu
        order_assumed = True
    else:
        order_assumed = False
    if not blocks:
        raise ValueError(f'{path.name}: nenalezen sloupec percentilu. Hlavička: {headers}')
    sha = sha256_of(path)
    records, unique = [], set()
    for row in rows[2:]:
        if not row:
            continue
        redizo = re.sub(r'\.0$', '', str(row[redizo_i] or '').strip())
        if not re.fullmatch(r'\d{9}', redizo):
            continue
        unique.add(redizo)
        rec = {
            'jpzYear': year, 'redizo': redizo,
            'oborovaSkupina': str(row[group_i]) if group_i is not None and row[group_i] is not None else None,
            'rocnik': str(row[rocnik_i]) if rocnik_i is not None and row[rocnik_i] is not None else None,
            'kraj': str(row[kraj_i]) if kraj_i is not None and row[kraj_i] is not None else None,
            'sourceFileSha256': sha, 'sourceUrl': source_url,
        }
        for b in blocks:
            rec[f'{b["subject"]}_took'] = as_count(row[b['took']]) if b['took'] is not None else None
            rec[f'{b["subject"]}_percentile'] = as_number(row[b['percentile']])
            rec[f'{b["subject"]}_sd'] = as_number(row[b['sd']]) if b['sd'] is not None else None
        records.append(rec)
    manifest = {'file': path.name, 'url': source_url, 'sha256': sha, 'bytes': path.stat().st_size,
                'physical_rows': len(rows), 'school_rows': len(records), 'unique_redizo': len(unique),
                'blocks': blocks, 'subject_order_assumed': order_assumed,
                'headers': [f'{g} | {l}' for g, l in headers]}
    return records, manifest


def is_four_year(text) -> bool:
    t = strip_accents(text or '')
    return bool(re.search(r'4|CTYR', t)) and not re.search(r'6|8|SEST|OSM|NASTAV', t)


def intake_by_school(jpz: list[dict]) -> dict[tuple[str, int], dict]:
    """Vstupní signál školy: percentil konajících uchazečů čtyřletých oborů, vážený počtem konajících."""
    acc = defaultdict(lambda: {'weight': 0.0, 'sum': 0.0, 'kraj': None})
    for r in jpz:
        if r['rocnik'] is not None and not is_four_year(r['rocnik']):
            continue
        for subject in ('cj', 'ma'):
            n, p = r.get(f'{subject}_took'), r.get(f'{subject}_percentile')
            if n and p is not None:
                key = (r['redizo'], r['jpzYear'])
                acc[key]['weight'] += n
                acc[key]['sum'] += n * p
                acc[key]['kraj'] = acc[key]['kraj'] or r['kraj']
    return {k: {'intakePercentile': v['sum'] / v['weight'], 'intakeN': v['weight'], 'kraj': v['kraj']}
            for k, v in acc.items() if v['weight'] > 0}


# ---------------------------------------------------------------------------
# Vrstva 2: kontextové srovnání a stabilita
# ---------------------------------------------------------------------------

def qualifier(mean, sd, n, reference):
    """Kvalifikátor ze standardní chyby průměru (§5.2)."""
    if mean is None or reference is None:
        return None, None, None
    if sd is None or not n:
        return None, None, 'chybí SD nebo n'
    se = sd / math.sqrt(n)
    low, high = mean - Z_95 * se, mean + Z_95 * se
    if low > reference:
        return 'nad_skupinou', round(se, 3), None
    if high < reference:
        return 'pod_skupinou', round(se, 3), None
    return 'nerozlisitelne', round(se, 3), None


def context_comparison(records: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for r in records:
        if r['grain'] == 'redizo_smo16' and r['averagePercentile'] is not None and (r['took'] or 0) >= MIN_N_REFERENCE:
            groups[(r['schoolYear'], r['period'], r['smo16'], r['subject'])].append(r)
    reference = {}
    for key, rs in groups.items():
        reference[key] = {
            'medianOfSchools': statistics.median(r['averagePercentile'] for r in rs),
            'pupilWeightedMean': sum(r['averagePercentile'] * r['took'] for r in rs) / sum(r['took'] for r in rs),
            'schools': len(rs), 'students': sum(r['took'] for r in rs),
        }
    output = []
    for r in records:
        if r['grain'] != 'redizo_smo16' or r['averagePercentile'] is None:
            continue
        ref = reference.get((r['schoolYear'], r['period'], r['smo16'], r['subject']))
        if not ref or ref['schools'] < 5:
            continue
        q, se, note = qualifier(r['averagePercentile'], r['standardDeviation'], r['took'], ref['medianOfSchools'])
        output.append({
            'id': r['id'], 'redizo': r['redizo'], 'smo16': r['smo16'], 'subject': r['subject'],
            'schoolYear': r['schoolYear'], 'period': r['period'], 'n': r['took'],
            'averagePercentile': r['averagePercentile'],
            'referenceMedian': round(ref['medianOfSchools'], 2), 'referenceSchools': ref['schools'],
            'referenceStudents': ref['students'],
            'differencePoints': round(r['averagePercentile'] - ref['medianOfSchools'], 2),
            'standardError': se, 'qualifier': q, 'note': note,
            'suppressed': (r['took'] or 0) < MIN_N_REFERENCE,
        })
    return output


def stability(context: list[dict]) -> list[dict]:
    series = defaultdict(list)
    for c in context:
        if c['period'] == 'jaro' and not c['suppressed']:
            series[(c['redizo'], c['smo16'], c['subject'])].append(c)
    output = []
    for (redizo, smo16, subject), items in series.items():
        items.sort(key=lambda c: c['schoolYear'])
        years = [c['schoolYear'] for c in items]
        last3 = [c for c in items if c['schoolYear'] >= years[-1] - 2]
        output.append({
            'redizo': redizo, 'smo16': smo16, 'subject': subject, 'years': years,
            'rolling3Mean': round(statistics.mean(c['averagePercentile'] for c in last3), 2),
            'rolling3Years': [c['schoolYear'] for c in last3],
            'yearsAbove': sum(c['qualifier'] == 'nad_skupinou' for c in items),
            'yearsBelow': sum(c['qualifier'] == 'pod_skupinou' for c in items),
            'yearsIndistinguishable': sum(c['qualifier'] == 'nerozlisitelne' for c in items),
            'breaks': sorted({c['schoolYear'] for c in items if c['schoolYear'] in COVID_MZ_YEARS or c['schoolYear'] == 2021}),
        })
    return output


# ---------------------------------------------------------------------------
# Vrstva 3: backtest (OLS bez externích knihoven)
# ---------------------------------------------------------------------------

def solve(a: list[list[float]], b: list[float]) -> list[float]:
    n = len(a)
    m = [row[:] + [b[i]] for i, row in enumerate(a)]
    for col in range(n):
        pivot = max(range(col, n), key=lambda r: abs(m[r][col]))
        if abs(m[pivot][col]) < 1e-12:
            raise ValueError('singulární matice (kolineární efekty)')
        m[col], m[pivot] = m[pivot], m[col]
        for r in range(n):
            if r != col:
                f = m[r][col] / m[col][col]
                for c in range(col, n + 1):
                    m[r][c] -= f * m[col][c]
    return [m[i][n] / m[i][i] for i in range(n)]


def ols(X: list[list[float]], y: list[float]) -> tuple[list[float], list[float], float]:
    k = len(X[0])
    xtx = [[sum(row[i] * row[j] for row in X) for j in range(k)] for i in range(k)]
    xty = [sum(row[i] * yi for row, yi in zip(X, y)) for i in range(k)]
    for i in range(k):
        xtx[i][i] += 1e-9
    beta = solve(xtx, xty)
    fitted = [sum(b * x for b, x in zip(beta, row)) for row in X]
    residuals = [yi - fi for yi, fi in zip(y, fitted)]
    mean_y = statistics.mean(y)
    ss_tot = sum((yi - mean_y) ** 2 for yi in y)
    ss_res = sum(r * r for r in residuals)
    return beta, residuals, (1 - ss_res / ss_tot if ss_tot else float('nan'))


def pearson(xs: list[float], ys: list[float]) -> float | None:
    if len(xs) < 3:
        return None
    mx, my = statistics.mean(xs), statistics.mean(ys)
    sxx = sum((x - mx) ** 2 for x in xs)
    syy = sum((y - my) ** 2 for y in ys)
    if sxx == 0 or syy == 0:
        return None
    return sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / math.sqrt(sxx * syy)


def build_pairs(mz: list[dict], intake: dict, four_year_smo16) -> list[dict]:
    pairs = []
    for r in mz:
        if r['grain'] != 'redizo_smo16' or r['period'] != 'jaro' or r['subject'] != 'cj':
            continue
        if r['averagePercentile'] is None or (r['took'] or 0) < MIN_N_BACKTEST:
            continue
        if four_year_smo16 is not None and r['smo16'] not in four_year_smo16:
            continue
        if four_year_smo16 is None and r['smo16Label'] and not is_four_year(r['smo16Label']):
            continue
        key = (r['redizo'], r['schoolYear'] - COHORT_SHIFT)
        if key not in intake:
            continue
        pairs.append({
            'redizo': r['redizo'], 'smo16': r['smo16'], 'mzYear': r['schoolYear'], 'jpzYear': key[1],
            'kraj': r['kraj'] or intake[key]['kraj'] or 'x', 'n': r['took'],
            'outcome': r['averagePercentile'], 'intake': intake[key]['intakePercentile'],
            'intakeN': intake[key]['intakeN'],
        })
    return pairs


def run_specification(pairs: list[dict], name: str, min_n: int, exclude_covid: bool) -> dict:
    data = [p for p in pairs if p['n'] >= min_n]
    if exclude_covid:
        data = [p for p in data if p['mzYear'] not in COVID_MZ_YEARS and p['jpzYear'] not in COVID_JPZ_YEARS]
    if len(data) < 30:
        return {'name': name, 'pairs': len(data), 'status': 'nedostatek párů'}
    smo_levels = sorted({p['smo16'] for p in data})[1:]
    kraj_levels = sorted({p['kraj'] for p in data})[1:]
    year_levels = sorted({p['mzYear'] for p in data})[1:]
    X, y = [], []
    for p in data:
        row = [1.0, p['intake'], math.log(p['n'])]
        row += [1.0 if p['smo16'] == s else 0.0 for s in smo_levels]
        row += [1.0 if p['kraj'] == k else 0.0 for k in kraj_levels]
        row += [1.0 if p['mzYear'] == yv else 0.0 for yv in year_levels]
        X.append(row)
        y.append(p['outcome'])
    try:
        beta, residuals, r2 = ols(X, y)
    except ValueError as exc:
        return {'name': name, 'pairs': len(data), 'status': str(exc)}
    by_school_year = {(p['redizo'], p['smo16'], p['mzYear']): res for p, res in zip(data, residuals)}
    lag_x, lag_y = [], []
    for (redizo, smo16, year), res in by_school_year.items():
        nxt = by_school_year.get((redizo, smo16, year + 1))
        if nxt is not None:
            lag_x.append(res)
            lag_y.append(nxt)
    r_lag1 = pearson(lag_x, lag_y)
    schools = defaultdict(list)
    for p, res in zip(data, residuals):
        schools[(p['redizo'], p['smo16'])].append(res)
    multi = [v for v in schools.values() if len(v) >= 2]
    within = statistics.mean(statistics.pvariance(v) for v in multi) if multi else None
    between = statistics.pvariance([statistics.mean(v) for v in multi]) if len(multi) > 1 else None
    return {
        'name': name, 'pairs': len(data), 'schools': len(schools), 'status': 'ok',
        'r2': round(r2, 3), 'intakeCoefficient': round(beta[1], 3), 'logNCoefficient': round(beta[2], 3),
        'residualSd': round(statistics.pstdev(residuals), 3),
        'lag1Pairs': len(lag_x), 'lag1Correlation': None if r_lag1 is None else round(r_lag1, 3),
        'betweenSchoolVariance': None if between is None else round(between, 3),
        'withinSchoolVariance': None if within is None else round(within, 3),
    }


def backtest(mz: list[dict], jpz: list[dict], four_year_smo16=None) -> dict:
    intake = intake_by_school(jpz)
    pairs = build_pairs(mz, intake, four_year_smo16)
    specs = [
        run_specification(pairs, 'zakladni', MIN_N_BACKTEST, False),
        run_specification(pairs, 'bez_covid_rocniku', MIN_N_BACKTEST, True),
        run_specification(pairs, 'n_alespon_30', 30, False),
    ]
    valid = [s for s in specs if s.get('status') == 'ok' and s.get('lag1Correlation') is not None]
    if not valid:
        decision = 'nelze_rozhodnout'
    elif all(s['lag1Correlation'] >= PUBLISH_THRESHOLD_R for s in valid):
        decision = 'predlozit_metodicke_oponenture'
    else:
        decision = 'nepublikovat'
    return {
        'criterion': f'korelace reziduí školy mezi po sobě jdoucími ročníky ≥ {PUBLISH_THRESHOLD_R} ve všech specifikacích (§5.3)',
        'cohortShiftYears': COHORT_SHIFT, 'intakeDefinition': 'percentil konajících uchazečů čtyřletých oborů na škole, vážený počtem konajících; nikoli přijatí',
        'pairsTotal': len(pairs), 'yearPairs': sorted({(p['jpzYear'], p['mzYear']) for p in pairs}),
        'specifications': specs, 'decision': decision,
        'limits': [
            'vstup popisuje uchazeče konající na škole, ne přijaté; u přeplněných škol je vstup podhodnocen (§5.3)',
            'bez proměnné přeplněnosti pro roky před 2024',
            'agregáty, ne stejní žáci; opakování, přestupy a organizační změny nejsou zachyceny',
            'model bez fixního efektu školy; korelace reziduí je diagnostika, nikoli odhad přidané hodnoty',
        ],
    }


# ---------------------------------------------------------------------------
# Syntetická data pro samotest
# ---------------------------------------------------------------------------

def synthetic_world(seed: int, school_effect_sd: float, schools: int = 240):
    rng = random.Random(seed)
    smo = {1: 'Gymnázium 4leté', 3: 'Lyceum', 5: 'SOŠ technické 1', 7: 'SOŠ ekonomické'}
    world = []
    for i in range(schools):
        s = rng.choice(list(smo))
        base = {1: 75, 3: 60, 5: 45, 7: 50}[s]
        world.append({
            'redizo': f'6{i:08d}', 'smo16': s, 'label': smo[s], 'kraj': f'CZ0{rng.randint(10, 80)}',
            'intake': max(5, min(95, rng.gauss(base, 10))), 'effect': rng.gauss(0, school_effect_sd),
            'n': rng.choice([8, 15, 25, 40, 60, 90]),
        })
    return world, rng


def write_synthetic_files(directory: Path, seed: int = 7, school_effect_sd: float = 4.0) -> None:
    """Zapíše XLSX ve tvaru, který loader očekává. Test loaderu je tím pádem kruhový; ověřuje běh, ne CERMAT."""
    world, rng = synthetic_world(seed, school_effect_sd)
    directory.mkdir(parents=True, exist_ok=True)
    for jpz_year in range(2017, 2023):
        book = openpyxl.Workbook()
        ws = book.active
        ws.title = f'JPZ{jpz_year}_red'
        ws.append([None] * 8 + ['ČESKÝ JAZYK'] + [None] * 4 + ['MATEMATIKA'] + [None] * 4)
        ws.append(['REDIZO / KRAJ / OBOROVÁ SKUPINA', 'OBOROVÁ SKUPINA', 'ROČNÍK', 'NÁZEV ŠKOLY', 'ADRESA ŠKOLY',
                   'KRAJ (KÓD)', 'KRAJ (NÁZEV)', 'ZŘIZOVATEL',
                   'PŘIHLÁŠENI', 'KONALI', 'NEKONALI', 'PRŮMĚRNÉ PERCENTILOVÉ UMÍSTĚNÍ', 'SMĚRODATNÁ ODCHYLKA (PERCENTIL. UMÍSTĚNÍ)',
                   'PŘIHLÁŠENI', 'KONALI', 'NEKONALI', 'PRŮMĚRNÉ PERCENTILOVÉ UMÍSTĚNÍ', 'SMĚRODATNÁ ODCHYLKA (PERCENTIL. UMÍSTĚNÍ)'])
        for w in world:
            drift = rng.gauss(0, 2)
            n = int(w['n'] * 1.8)
            ws.append([w['redizo'], w['label'], '4leté', 'Škola', 'Adresa', w['kraj'], 'Kraj', 'kraj',
                       n + 3, n, 3, round(w['intake'] + drift, 2), 20.0,
                       n + 3, n, 3, round(w['intake'] + drift - 3, 2), 21.0])
        book.save(directory / f'JPZ{jpz_year}_skoly-skolobory_vysledky.xlsx')
    for mz_year in range(2021, 2027):
        book = openpyxl.Workbook()
        ws = book.active
        ws.title = str(mz_year)
        top = ['', '', '', '', '', '', 'SPOLEČNÁ ČÁST', None, None, None, None, None,
               'ČESKÝ JAZYK A LITERATURA', None, None, None, None, None,
               'MATEMATIKA', None, None, None, None, None, None]
        bottom = ['REDIZO', 'NÁZEV ŠKOLY', 'ÚROVEŇ AGREGACE', 'SMO16', 'SMO16 - NÁZEV', 'KRAJ',
                  'PŘIHLÁŠENI', 'KONALI', 'NEKONALI', 'USPĚLI', 'NEUSPĚLI', 'ČISTÁ NEÚSPĚŠNOST (%)',
                  'KONALI', 'USPĚLI', 'NEUSPĚLI', 'PRŮMĚRNÝ % SKÓR', 'SMĚRODATNÁ ODCHYLKA', 'PRŮMĚRNÝ PERCENTIL',
                  'PODÍL VOLBY (%)', 'KONALI', 'USPĚLI', 'NEUSPĚLI', 'PRŮMĚRNÝ % SKÓR', 'SMĚRODATNÁ ODCHYLKA', 'PRŮMĚRNÝ PERCENTIL']
        ws.append(top)
        ws.append(bottom)
        ws.append(['', 'celkem', 'total', '', '', '', 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0])
        for w in world:
            n = w['n']
            cj = max(3, min(97, 0.8 * w['intake'] + 10 + w['effect'] + rng.gauss(0, 14 / math.sqrt(n))))
            sd = 18.0 + rng.gauss(0, 2)
            failed = int(round(n * max(0.0, (60 - cj) / 200)))
            ma_share = max(0.1, min(0.9, (w['intake'] - 20) / 80))
            ma_n = max(1, int(round(n * ma_share)))
            ma = max(3, min(97, cj - 5 + rng.gauss(0, 20 / math.sqrt(ma_n))))
            row = [w['redizo'], 'Škola', 'redizo_smo16', w['smo16'], w['label'], w['kraj'],
                   n + 2, n, 2, n - failed, failed, round(100 * failed / n, 2),
                   n, n - failed, failed, round(cj * 0.9, 2), round(sd, 2), round(cj, 2),
                   round(100 * ma_share, 1), ma_n, ma_n, 0, round(ma * 0.9, 2), 21.0, round(ma, 2)]
            ws.append(row)
            ws.append([w['redizo'], 'Škola', 'redizo', '', '', w['kraj']] + row[6:])
        vys = book.create_sheet('vysvetlivky')
        vys['A2'], vys['B2'] = 'Populace', 'prvomaturanti za daný rok celkem (SYNTETICKÁ DATA)'
        book.save(directory / f'MZ{mz_year}j_SC_skolobory.xlsx')


# ---------------------------------------------------------------------------
# Hlavní běh
# ---------------------------------------------------------------------------

def run(input_dir: Path, lenient: bool = False, four_year_smo16=None) -> dict:
    mz_records, jpz_records, sources, problems = [], [], [], []
    for year in MZ_YEARS:
        for suffix, period in (('j', 'jaro'), ('jap', 'stav_po_podzimu')):
            name = f'MZ{year}{suffix}_SC_skolobory.xlsx'
            path = input_dir / name
            if not path.exists():
                continue
            try:
                recs, manifest = load_mz(path, year, period, SOURCE_BASE_MZ + name, lenient)
            except ValueError as exc:
                problems.append(str(exc))
                continue
            mz_records += recs
            sources.append(manifest)
    for year in JPZ_YEARS:
        name = f'JPZ{year}_skoly-skolobory_vysledky.xlsx'
        path = input_dir / name
        if not path.exists():
            continue
        try:
            recs, manifest = load_jpz(path, year, SOURCE_BASE_JPZ + name)
        except ValueError as exc:
            problems.append(str(exc))
            continue
        jpz_records += recs
        sources.append(manifest)
    context = context_comparison(mz_records)
    counts, by_year, jpz_by_year = defaultdict(int), defaultdict(int), defaultdict(int)
    for c in context:
        counts[c['qualifier'] or 'none'] += 1
    for r in mz_records:
        by_year[f'{r["schoolYear"]}_{r["period"]}_{r["grain"]}'] += 1
    for r in jpz_records:
        jpz_by_year[str(r['jpzYear'])] += 1
    return {
        'generated_at': datetime.now(timezone.utc).isoformat(),
        'code_commit': git_commit(),
        'input_dir': str(input_dir),
        'sources': sources, 'problems': problems,
        'coverage': {
            'mz_records': len(mz_records),
            'mz_by_year_period_grain': dict(sorted(by_year.items())),
            'jpz_rows': len(jpz_records),
            'jpz_rows_by_year': dict(sorted(jpz_by_year.items())),
            'context_rows': len(context), 'qualifier_counts': dict(counts),
        },
        'records': mz_records,
        'context': context,
        'stability': stability(context),
        'backtest': backtest(mz_records, jpz_records, four_year_smo16),
        'limits': [
            'společná část = didaktické testy; nikoli profilová část ani celá maturita',
            'redizo_smo16 není KKOV; oborový výsledek jen při jednoznačné mapě',
            'mapa sloupců MZ vznikla bez přístupu ke skutečnému souboru; ověřit proti listu vysvětlivky',
            'výstup backtestu je vstup pro oponenturu, nikoli důkaz přidané hodnoty',
        ],
    }


def self_test(scratch: Path, output: Path | None = None) -> None:
    # 1. kvalifikátor
    assert qualifier(60, 20, 100, 55)[0] == 'nad_skupinou'
    assert qualifier(59, 20, 18, 55)[0] == 'nerozlisitelne'   # +4 body při n=18 je šum
    assert qualifier(40, 20, 100, 55)[0] == 'pod_skupinou'
    assert qualifier(60, None, 100, 55)[0] is None
    # 2. OLS na známé funkci
    beta, _, r2 = ols([[1, x] for x in range(10)], [3 + 2 * x for x in range(10)])
    assert abs(beta[0] - 3) < 1e-6 and abs(beta[1] - 2) < 1e-6 and r2 > 0.999
    # 3. svět se stabilním školním efektem: rezidua musí být korelovaná
    stable_dir = scratch / 'stable'
    write_synthetic_files(stable_dir, seed=7, school_effect_sd=4.0)
    result = run(stable_dir)
    assert not result['problems'], result['problems']
    assert result['coverage']['mz_records'] > 0 and result['coverage']['jpz_rows'] > 0
    base = result['backtest']['specifications'][0]
    assert base['status'] == 'ok', base
    assert base['lag1Correlation'] is not None and base['lag1Correlation'] > 0.4, base
    assert result['backtest']['decision'] == 'predlozit_metodicke_oponenture', result['backtest']
    qc = result['coverage']['qualifier_counts']
    assert qc.get('nerozlisitelne', 0) > 0 and (qc.get('nad_skupinou', 0) + qc.get('pod_skupinou', 0)) > 0, qc
    # 4. svět bez školního efektu: korelace reziduí musí spadnout k nule
    null_dir = scratch / 'null'
    write_synthetic_files(null_dir, seed=11, school_effect_sd=0.0)
    null_result = run(null_dir)
    null_base = null_result['backtest']['specifications'][0]
    assert null_base['status'] == 'ok' and abs(null_base['lag1Correlation']) < 0.2, null_base
    assert null_result['backtest']['decision'] == 'nepublikovat', null_result['backtest']['decision']
    # 5. kontrakt záznamu
    sample = result['records'][0]
    for field in ('schoolYear', 'period', 'grain', 'redizo', 'smo16', 'subject', 'took', 'averagePercentile',
                  'standardDeviation', 'populationNote', 'sourceUrl', 'sourceFileSha256', 'quality', 'qualityReasons'):
        assert field in sample, field
    summary = {
        'self_test': 'ok',
        'checked_at': datetime.now(timezone.utc).isoformat(),
        'code_commit': git_commit(),
        'data': 'syntetická; datový server CERMAT nebyl z vývojového prostředí dostupný',
        'design': {
            'stable_world': 'školní efekt SD 4 percentilové body, stálý v čase',
            'null_world': 'školní efekt SD 0; veškerý rozptyl je vstup a šum',
            'outcome_model': 'MZ ČJ percentil = 0,8 · vstup + 10 + školní efekt + šum/sqrt(n)',
        },
        'stable_world': {k: base[k] for k in ('pairs', 'schools', 'r2', 'intakeCoefficient', 'residualSd',
                                              'lag1Pairs', 'lag1Correlation', 'betweenSchoolVariance', 'withinSchoolVariance')},
        'stable_decision': result['backtest']['decision'],
        'null_world': {k: null_base[k] for k in ('pairs', 'schools', 'r2', 'intakeCoefficient', 'residualSd',
                                                 'lag1Pairs', 'lag1Correlation', 'betweenSchoolVariance', 'withinSchoolVariance')},
        'null_decision': null_result['backtest']['decision'],
        'publish_threshold': PUBLISH_THRESHOLD_R,
        'qualifier_counts': qc,
        'limits': [
            'syntetická data potvrzují, že kritérium §5.3 rozliší stabilní efekt od šumu; neříkají nic o skutečných školách',
            'loader byl testován proti souborům, které skript sám zapsal; mapa sloupců MZ zůstává neověřená proti CERMAT',
        ],
    }
    print(json.dumps(summary, ensure_ascii=False, indent=1))
    if output:
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(summary, ensure_ascii=False, indent=1))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--input-dir', type=Path, help='adresář se staženými XLSX CERMAT')
    parser.add_argument('--output', type=Path, help='cílový JSON (doporučeno docs/podklady/)')
    parser.add_argument('--self-test', action='store_true', help='běh na syntetických datech')
    parser.add_argument('--scratch', type=Path, default=Path('/tmp/maturita-prototyp'), help='adresář pro syntetická data')
    parser.add_argument('--lenient', action='store_true', help='nepřerušit při chybějících povinných sloupcích MZ')
    parser.add_argument('--four-year-smo16', type=int, nargs='*', help='kódy SMO16 čtyřletých oborů pro backtest; jinak podle názvu skupiny')
    parser.add_argument('--no-records', action='store_true', help='nezapisovat dlouhou tabulku záznamů do výstupu')
    args = parser.parse_args()
    if args.self_test:
        self_test(args.scratch, args.output)
        return
    if not args.input_dir or not args.output:
        parser.error('zadejte --input-dir a --output, nebo --self-test')
    result = run(args.input_dir, args.lenient, set(args.four_year_smo16) if args.four_year_smo16 else None)
    if args.no_records:
        result['records'] = f'{len(result["records"])} záznamů vynecháno (--no-records)'
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=1))
    summary = {k: result[k] for k in ('coverage', 'problems')}
    summary['backtest'] = {k: result['backtest'][k] for k in ('pairsTotal', 'decision', 'specifications')}
    print(json.dumps(summary, ensure_ascii=False, indent=1))
    if result['problems']:
        sys.exit(2)


if __name__ == '__main__':
    main()
