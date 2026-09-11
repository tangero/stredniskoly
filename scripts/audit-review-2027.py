"""Reprodukovatelné podklady oponentury; nemění produkční exporty.

python3 scripts/audit-review-2027.py --cache /tmp/gymnazium-oponentura-r1 \
  --output docs/podklady/oponentura-2027-r1.json
Vyžaduje openpyxl. Chybějící XLSX stáhne z CERMAT; SHA-256 zachytí revizi.
"""
import argparse
import hashlib
import json
import re
import subprocess
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import urlopen

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def normalize(key: str) -> str:
    parts = key.split('_')
    text = unicodedata.normalize('NFKD', '_'.join(parts[2:]))
    text = ''.join(c for c in text if not unicodedata.combining(c))
    suffix = re.sub('[^a-zA-Z0-9]+', '_', text).strip('_').lower()
    return '_'.join(parts[:2]) + ('_' + suffix if suffix else '')


def base(key: str) -> str:
    return '_'.join(key.split('_')[:2])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cache', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    args.cache.mkdir(parents=True, exist_ok=True)
    urls = [f'https://data.cermat.cz/files/files/JPZ/agregovana_data_skoly/JPZ{y}_skoly-skolobory_vysledky.xlsx' for y in range(2017, 2024)]
    urls += [f'https://data.cermat.cz/files/files/MZ/agregovana_data_skoly/MZ{period}_SC_skolobory.xlsx' for period in ['2026j', '2025jap']]
    urls += ['https://data.cermat.cz/files/files/JPZ/agregovana_data_polozky/2026/JPZ2026_M4_ulohy_agregovane_vysledky.xlsx']
    sources = []
    for url in urls:
        path = args.cache / url.rsplit('/', 1)[1]
        if not path.exists():
            with urlopen(url, timeout=45) as response:
                path.write_bytes(response.read())
        data = path.read_bytes()
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = list(book[book.sheetnames[0]].values)
        item = {'url': url, 'sha256': hashlib.sha256(data).hexdigest(), 'bytes': len(data),
                'sheets': book.sheetnames, 'physical_rows': len(rows), 'columns': max(map(len, rows))}
        if path.name.startswith('MZ'):
            item['aggregation_rows'] = dict(Counter(r[2] for r in rows[2:]))
            item['population'] = book['vysvetlivky']['B2'].value
        elif 'ulohy' in path.name:
            item['sheet_scope'] = rows[1][4]
            item['participants'] = rows[3][3]
            item['mean_percent_score'] = rows[3][12]
            item['item_01_3_mean_percent_score'] = next(r[2] for r in rows if r[1] == '01.3')
        else:
            schools = [r for r in rows[2:] if re.fullmatch(r'\d{9}', str(r[0]))]
            item['school_rows'] = len(schools)
            item['unique_redizo'] = len({r[0] for r in schools})
            item['header'] = [x for x in rows[1] if x is not None]
        book.close()
        sources.append(item)

    old_path, new_path = ROOT / 'public/schools_data.json', ROOT / 'public/cermat_results_2026.json'
    old = json.loads(old_path.read_text())['2025']
    new = json.loads(new_path.read_text())
    target_groups = defaultdict(list)
    for key, value in new.items():
        target_groups[base(key)].append({'id': key, 'mean': value['cj_ma_prijati'], 'admitted': value['prijati']})
    target_norm = Counter(normalize(k) for k in new)
    source_norm = Counter(normalize(r['id']) for r in old)
    source_groups = Counter(base(r['id']) for r in old)
    methods = {}
    for label, ids in [('rows_2025', [r['id'] for r in old]), ('unique_ids_2025', sorted({r['id'] for r in old}))]:
        methods[label] = {
            'denominator': len(ids), 'exact': sum(k in new for k in ids),
            'base_exists': sum(base(k) in target_groups for k in ids),
            'base_multiple_targets': sum(len(target_groups.get(base(k), [])) > 1 for k in ids),
            'base_unique_both_raw': sum(len(target_groups.get(base(k), [])) == 1 and source_groups[base(k)] == 1 for k in ids),
            'normalized_unique_both': sum(target_norm[normalize(k)] == 1 and source_norm[normalize(k)] == 1 for k in ids),
        }
    examples = [{'base': key, 'source_ids': [r['id'] for r in old if base(r['id']) == key], 'targets': rows}
                for key, rows in target_groups.items() if len(rows) > 1 and source_groups[key]][:8]
    output = {
        'checked_at': datetime.now(timezone.utc).isoformat(),
        'code_commit': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(),
        'source_files': sources,
        'local_sha256': {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in [old_path, new_path]},
        'pairing': {'methods': methods, 'target_collision_groups': sum(len(v) > 1 for v in target_groups.values()), 'counterexamples': examples},
        'limits': ['Počet existujících základních klíčů není počet správně spárovaných zaměření.',
                   'Fyzické řádky XLSX zahrnují hlavičky a souhrny.',
                   'Inventura a kontrola vzorků nejsou produkční import ani validace kohortního modelu.'],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(output['pairing']['methods'], ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
