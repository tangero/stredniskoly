#!/usr/bin/env python3
"""Historie jedné nabídky napříč všemi dostupnými koly přijímacího řízení.

Fronta návazností srovnává jen 1. kolo 2025 a 2026. Repozitář ale obsahuje
i další soubory CERMAT, které umožňují odpovědět na dvě opakující se otázky
rešerše: běžel obor i před rokem 2025, a byl nabízen mimo 1. kolo?

Pokryté zdroje:

| Rok a kolo | Soubor | Klíč |
|---|---|---|
| 2024, 1. kolo | PZ2024_kolo1_uchazeci_prihlasky_vysledky.xlsx | IZO |
| 2024, 2. kolo | PZ2024_kolo2_uchazeci_prihlasky_vysledky.xlsx | IZO |
| 2025, 1. kolo | PZ2025_kolo1_uchazeci_prihlasky_vysledky.xlsx | REDIZO |
| 2025, 2. kolo | PZ2025_kolo2_uchazeci_prihlasky_vysledky.xlsx | REDIZO |
| 2026, 1. kolo | PZ2026_kolo1_skolobory_kapacity.xlsx | REDIZO i IZO |

Soubory z let 2024 a 2025 jsou na úrovni uchazečů: obor bez jediné přihlášky
v nich chybí, i kdyby vypsán byl. Přítomnost je proto silný důkaz, nepřítomnost
slabý. Data 2. kola 2026 zveřejněna nejsou, tuto otázku nelze uzavřít.

Použití:
    python3 scripts/offer-history.py --redizo 600007022 --izo 107820102 --kkov 64-41-L/51
    python3 scripts/offer-history.py --task R-600007022        # všechny nabídky úkolu
"""
import argparse
import json
from collections import defaultdict
from pathlib import Path

import openpyxl

DATA = Path("data")
FRONTA = Path("docs/podklady/fronta-dohledavani-2025-2026.json")

ZDROJE_2024 = [("2024, 1. kolo", "PZ2024_kolo1_uchazeci_prihlasky_vysledky.xlsx"),
               ("2024, 2. kolo", "PZ2024_kolo2_uchazeci_prihlasky_vysledky.xlsx")]
ZDROJE_2025 = [("2025, 1. kolo", "PZ2025_kolo1_uchazeci_prihlasky_vysledky.xlsx"),
               ("2025, 2. kolo", "PZ2025_kolo2_uchazeci_prihlasky_vysledky.xlsx")]
ZDROJ_2026 = ("2026, 1. kolo", "PZ2026_kolo1_skolobory_kapacity.xlsx")

FORMY_DENNI = {"den", "den2"}


def bez_prefixu(izo):
    t = str(izo or "")
    return t[4:] if t.lower().startswith("izo_") else t


def nacti_2024(cesta):
    """{(IZO, KKOV): (přihlášek, přijatých)} – sloupce po školách, hodnoty True/False."""
    wb = openpyxl.load_workbook(cesta, read_only=True)
    ws = wb.worksheets[0]
    prihlasky, prijati = defaultdict(int), defaultdict(int)
    for r in ws.iter_rows(min_row=2, values_only=True):
        for i in range(5):
            z = 5 + i * 7
            izo, kkov, forma, zkraceno, prijat = r[z], r[z + 2], r[z + 3], r[z + 4], r[z + 5]
            if not izo or not kkov or forma not in FORMY_DENNI or str(zkraceno) == "True":
                continue
            klic = (str(izo).strip(), str(kkov).strip())
            prihlasky[klic] += 1
            if str(prijat) == "True":
                prijati[klic] += 1
    wb.close()
    return {k: (prihlasky[k], prijati[k]) for k in prihlasky}


def nacti_2025(cesta):
    """{(REDIZO, KKOV): (přihlášek, přijatých)} – sloupce po atributech, kódy 1/2."""
    wb = openpyxl.load_workbook(cesta, read_only=True)
    ws = wb.worksheets[0]
    prihlasky, prijati = defaultdict(int), defaultdict(int)
    for r in ws.iter_rows(min_row=2, values_only=True):
        for i in range(5):
            redizo, kkov, forma, zkraceno, prijat = r[2 + i], r[12 + i], r[17 + i], r[22 + i], r[27 + i]
            if not redizo or not kkov or forma not in FORMY_DENNI or zkraceno == 1:
                continue
            klic = (str(redizo).strip(), str(kkov).strip())
            prihlasky[klic] += 1
            if prijat == 1:
                prijati[klic] += 1
    wb.close()
    return {k: (prihlasky[k], prijati[k]) for k in prihlasky}


def nacti_2026(cesta):
    """{(REDIZO, KKOV): kapacita} a {(IZO, KKOV): kapacita} z přehledu škol a oborů."""
    wb = openpyxl.load_workbook(cesta, read_only=True)
    ws = wb.worksheets[0]
    hlavicka = [str(h) for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    idx = {n: hlavicka.index(n) for n in ("IZO", "REDIZO", "KKOV") if n in hlavicka}
    sloupec_kap = next((i for i, h in enumerate(hlavicka) if "KAPACIT" in h.upper()), None)
    podle = defaultdict(int)
    for r in ws.iter_rows(min_row=2, values_only=True):
        kkov = str(r[idx["KKOV"]]).strip() if r[idx["KKOV"]] else None
        if not kkov:
            continue
        kap = r[sloupec_kap] if sloupec_kap is not None else None
        kap = kap if isinstance(kap, (int, float)) else 0
        podle[(str(r[idx["REDIZO"]]).strip(), kkov)] += kap
        podle[(bez_prefixu(r[idx["IZO"]]), kkov)] += kap
    wb.close()
    return dict(podle)


def historie(redizo, izo, kkov, zdroje):
    radky = []
    for popis, mapa, klic in (
        [(p, zdroje[p], (izo, kkov)) for p, _ in ZDROJE_2024]
        + [(p, zdroje[p], (redizo, kkov)) for p, _ in ZDROJE_2025]
    ):
        hodnota = mapa.get(klic)
        radky.append((popis, f"{hodnota[0]} přihlášek, {hodnota[1]} přijatých" if hodnota else "není"))
    kap = zdroje[ZDROJ_2026[0]].get((redizo, kkov)) or zdroje[ZDROJ_2026[0]].get((izo, kkov))
    radky.append((ZDROJ_2026[0], f"kapacita {int(kap)}" if kap else "není"))
    return radky


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--redizo")
    parser.add_argument("--izo")
    parser.add_argument("--kkov")
    parser.add_argument("--task")
    args = parser.parse_args()

    zdroje = {}
    for popis, soubor in ZDROJE_2024:
        zdroje[popis] = nacti_2024(DATA / soubor)
    for popis, soubor in ZDROJE_2025:
        zdroje[popis] = nacti_2025(DATA / soubor)
    zdroje[ZDROJ_2026[0]] = nacti_2026(DATA / ZDROJ_2026[1])

    dotazy = []
    if args.task:
        fronta = json.loads(FRONTA.read_text())
        ukol = next((t for t in fronta["tasks"] if t["id"] == args.task), None)
        if not ukol:
            raise SystemExit(f"Úkol {args.task} není ve frontě.")
        videno = set()
        for rok in ("2025", "2026"):
            for n in (ukol.get("context_offers", {}).get(rok) or []):
                k = (str(n["REDIZO"]), bez_prefixu(n["IZO"]), n["KKOV"])
                if k not in videno:
                    videno.add(k)
                    dotazy.append(k)
    else:
        if not args.kkov or not (args.redizo or args.izo):
            raise SystemExit(__doc__)
        dotazy.append((args.redizo or "", bez_prefixu(args.izo), args.kkov))

    for redizo, izo, kkov in dotazy:
        print(f"\nREDIZO {redizo} | IZO {izo} | obor {kkov}")
        for popis, hodnota in historie(redizo, izo, kkov, zdroje):
            print(f"  {popis:16} {hodnota}")
    print("\nPozn.: 2024 a 2025 jsou data uchazečů (obor bez přihlášek chybí); "
          "2. kolo 2026 zveřejněno není.")


if __name__ == "__main__":
    raise SystemExit(main())
