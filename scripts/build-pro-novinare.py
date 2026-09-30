#!/usr/bin/env python3
"""Balíčky dat pro novináře: public/pro-novinare/.

Návrh: docs/navrh-pro-novinare-2027.md. Ukazatele: docs/slovnik-ukazatelu.md,
oddíl 10. Stránka /pro-novinare čte souhrn.json, soubory ke stažení leží vedle.

Šest balíčků, každý jako CSV (UTF-8, čárka, desetinná tečka, pro strojové
zpracování) a XLSX (s listem „O datech“, pro Excel):

  veletrhy             akce s potvrzeným termínem a počty po krajích
  konzervatore         konzervatoře z rejstříku MŠMT a termíny, přihlášky do 30. 11.
  obory-1-kolo-{rok}   nabídky 1. kola s poptávkou, výsledkem a stavem 2. kola
  uchazeci-{rok}        kam se uchazeči dostali v 1. a 2. kole, volná místa po 2. kole
  druhe-kolo-{rok}     všechny nabídky 2. kola a souhrny po krajích a typech škol
  kriteria-{rok}       co vedle jednotné zkoušky bodovalo, podle přepisu kritérií

Roky se berou z registru public/stav_datovych_sad.json, ne z kódu. Vstupní
soubory CERMAT se ověřují otiskem sha256 proti registru nebo metadatům webu;
jiný soubor, než jaký web převzal, generátor odmítne.

    python3 scripts/build-pro-novinare.py --vstupy ADRESAR
    python3 scripts/build-pro-novinare.py --vstupy ADRESAR --dnes 2026-09-29

ADRESAR obsahuje PZ{rok}_kolo1_skolobory_vysledky.xlsx,
PZ{rok}_kolo2_skolobory_vysledky.xlsx, PZ{rok}_kolo1_uchazeci_prihlasky_vysledky.xlsx
a PZ{rok}_kolo2_uchazeci_prihlasky_vysledky.xlsx
(zdrojové soubory nejsou v gitu, adresy jsou v registru).
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import importlib.util
import io
import json
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

KOREN = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(KOREN / "scripts"))
from slouceni_prihlasek import byl_prijat, denni_nezkracene  # noqa: E402

VYSTUP = KOREN / "public" / "pro-novinare"
REGISTR = KOREN / "public" / "stav_datovych_sad.json"
NASTAVENI = KOREN / "src" / "data" / "pro-novinare.json"

# Nástavbové studium (kategorie L5, například 64-41-L/51) je pro vyučené, ne pro žáky základní školy.
NASTAVBA = re.compile(r"^\d\d-\d\d-L/5")

# Stejné prahy jako na webu (docs/slovnik-ukazatelu.md): nejnižší výsledek
# přijatých se uvádí jen při aspoň deseti přijatých s výsledkem zkoušky.
MIN_PRIJATYCH_PRO_MINIMUM = 10

VYHRADA_KRITERII = (
    "Neověřený přepis PDF kritérií {rok}; zhruba každý desátý přepis obsahuje podstatnou chybu. "
    "Platí pro rok {rok}, ne pro nové přijímací řízení. Ověřte v kritériích školy."
)

# Skupiny složek kritérií podle názvu z přepisu. Jedna složka může patřit do
# více skupin („pohovor nad portfoliem“). Třídí se jen to, co bodovalo vedle
# jednotné zkoušky; minima a srážky za chování se netřídí.
SKUPINY_SLOZEK = {
    "prospech": re.compile(
        r"prospěch|prospech|vysvědčen|vysvedcen|průměr|prumer|známk|znamk|klasifikac|"
        r"výsledk\w* (z )?(dosavadního|předchozího|dosavadniho|predchoziho)|"
        r"(dosavadního|předchozího|dosavadniho|predchoziho) (studia|vzdělávání|vzdělání|vzdelani|vzdelavani)|"
        r"hodnocení (na|z|výsledků|předchozího|prospěchu)|\b\w?ZŠ\b|\bZS\b|studijní výsledky|"
        r"předchozí\w* vzděláv|výsledky vzdělávání|profilov\w* předmět|profilove predmety",
        re.I,
    ),
    "skolni_zkouska": re.compile(
        r"školní\w*\s+(přijímací\w*\s+)?(zkoušk|test)|skolni\w*\s+(prijimaci\w*\s+)?(zkousk|test)|"
        r"\bOSP\b|\bŠPZ\b|studijních předpokladů|písemn|ústní|didaktick\w* test školy|vlastní test|test školy|"
        r"test z |test při",
        re.I,
    ),
    "pohovor": re.compile(r"pohovor|rozhovor", re.I),
    "talentova_prakticka": re.compile(
        r"talentov|praktick|fyzick|tělesn|telesn|sportovní (zkoušk|výkon|testy)|pohybov|výtvarn|hudebn|"
        r"klauzur|domácí práce|domaci prace|portfoli",
        re.I,
    ),
    "souteze_aktivity": re.compile(
        r"soutěž|soutez|olympi|kroužk|krouzk|aktivit|mimoškol|mimoskol|zájm|zajm|certifik|"
        r"dalš\w* skutečnost|dalsi\w* skutecnost|klub|bonus|bonifikac|dobrovol|reprezent|ocenění|oceneni|"
        r"zájem|podnikavost|aktivní přístup|kroužc",
        re.I,
    ),
}
# Složky, které nejsou extra body: jednotná zkouška uvedená mezi složkami,
# součtové řádky a chování (srážky za chování do extra bodů nepatří,
# slovník pojmů, heslo „extra body“). Nezapočítají se do žádné skupiny.
NEJSOU_EXTRA_BODY = re.compile(
    r"státní přijímací|jednotn\w* (přijímací )?zkoušk|\bJPZ\b|\bJZ\b|didaktick\w* test\w*$|^celkem$|"
    r"počáteční stav|^(hodnocení )?chování",
    re.I,
)
NAZVY_SKUPIN = {
    "prospech": "prospěch ze základní školy",
    "skolni_zkouska": "školní přijímací zkouška",
    "pohovor": "pohovor",
    "talentova_prakticka": "talentová, praktická nebo sportovní zkouška, portfolio",
    "souteze_aktivity": "soutěže, certifikáty a další aktivity",
}


# ---------------------------------------------------------------------------
# Pomocné


def nacti_json(cesta: Path):
    return json.loads(cesta.read_text(encoding="utf-8"))


def sha256(cesta: Path) -> str:
    h = hashlib.sha256()
    with cesta.open("rb") as f:
        for blok in iter(lambda: f.read(1 << 20), b""):
            h.update(blok)
    return h.hexdigest()


def cislo(v) -> float | None:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def cele(v) -> int:
    c = cislo(v)
    return int(c) if c is not None else 0


def podil(citatel: float, jmenovatel: float, mista: int = 3) -> float | None:
    return round(citatel / jmenovatel, mista) if jmenovatel else None


def nadpis_kraje(nazev: str) -> str:
    """Jako nadpisKraje() v src/lib/kraje.mjs: „Středočeský kraj“, výjimky beze slova kraj."""
    nazev = (nazev or "").strip()
    if nazev in ("Hlavní město Praha", "Kraj Vysočina", KRAJ_NEURCEN) or not nazev:
        return nazev
    return f"{nazev} kraj"


def bez_relativniho_roku(text: str, rok_sezony: int) -> str:
    """„loni“ a „letos“ ze zdrojové poznámky nahradí rokem (slovník pojmů, pravidlo 4).

    Sezóna veletrhů je podzim roku `rok_sezony`; „loni“ je tedy rok před ním.
    """
    text = re.sub(r"\bloni\b", f"v roce {rok_sezony - 1}", text)
    text = re.sub(r"\bLoni\b", f"V roce {rok_sezony - 1}", text)
    text = re.sub(r"\bletošní ročník\b", f"ročník {rok_sezony}", text)
    text = re.sub(r"\bletošn(í|ího|ím)\b", f"z roku {rok_sezony}", text)
    text = re.sub(r"\b[Ll]etos\b", f"v roce {rok_sezony}", text)
    return text


KRAJ_NEURCEN = "kraj neurčen (škola má obor ve více krajích)"


def mapa_kraju(souhrn: list[dict]) -> dict[tuple[str, str], str]:
    """(REDIZO, KKOV) → název kraje ze souhrnu CERMAT.

    Data uchazečů nesou jen REDIZO a KKOV, ne zaměření ani místo. Škola může mít obory
    v různých krajích (PORG: 75-31-M/01 v Brně, 79-41-K/81 v Praze, Brně i Ostravě),
    proto se kraj bere podle dvojice; když ani ta není jednoznačná, vrací KRAJ_NEURCEN.
    """
    kraje: dict[tuple[str, str], set[str]] = defaultdict(set)
    for r in souhrn:
        kraje[(str(r["REDIZO"]), str(r["KKOV"]))].add(str(r["KRAJ - NÁZEV"]))
    return {k: (next(iter(v)) if len(v) == 1 else KRAJ_NEURCEN) for k, v in kraje.items()}


def nacti_xlsx(cesta: Path) -> list[dict]:
    wb = openpyxl.load_workbook(cesta, read_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    hlavicka = [str(h).strip() if h is not None else "" for h in next(it)]
    return [dict(zip(hlavicka, r)) for r in it]


def denni_nezkracene_radek(r: dict) -> bool:
    """Souhrny CERMAT: denní forma (den, den2) a nezkrácené studium, jako build-druhe-kolo.py."""
    return str(r.get("FORMA VZDĚLÁVÁNÍ") or "").startswith("den") and r.get("ZKRÁCENÉ STUDIUM") == "ne"


def nacti_druhe_kolo_modul():
    spec = importlib.util.spec_from_file_location("druhe_kolo", KOREN / "scripts" / "build-druhe-kolo.py")
    modul = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modul)
    return modul


# ---------------------------------------------------------------------------
# Zápis balíčků


class Balicek:
    """Jeden balíček: několik tabulek a text „O datech“. Zapisuje CSV za každou tabulku a jeden XLSX."""

    def __init__(self, klic: str, nazev: str, popis: str, o_datech: list[str]):
        self.klic = klic
        self.nazev = nazev
        self.popis = popis
        self.o_datech = o_datech
        self.tabulky: list[tuple[str, str, list[str], list[list]]] = []  # (soubor, list, hlavička, řádky)

    def pridej(self, soubor: str, list_: str, hlavicka: list[str], radky: list[list]):
        self.tabulky.append((soubor, list_, hlavicka, radky))

    def zapis(self, adresar: Path, citace: str) -> list[dict]:
        soubory = []
        for soubor, list_, hlavicka, radky in self.tabulky:
            cesta = adresar / f"{soubor}.csv"
            with cesta.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f, lineterminator="\n")
                w.writerow(hlavicka)
                w.writerows(radky)
            soubory.append({"soubor": cesta.name, "list": list_, "radku": len(radky), "format": "csv"})
        wb = openpyxl.Workbook()
        o = wb.active
        o.title = "O datech"
        o.append([self.nazev])
        o["A1"].font = Font(bold=True, size=14)
        o.append([])
        for veta in [self.popis, "", *self.o_datech, "", citace]:
            o.append([veta])
        o.column_dimensions["A"].width = 140
        for _, list_, hlavicka, radky in self.tabulky:
            ws = wb.create_sheet(list_[:31])
            ws.append(hlavicka)
            for bunka in ws[1]:
                bunka.font = Font(bold=True)
            for r in radky:
                ws.append(r)
            ws.freeze_panes = "A2"
            for i, h in enumerate(hlavicka, start=1):
                ws.column_dimensions[get_column_letter(i)].width = max(10, min(45, len(h) + 2))
        cesta = adresar / f"{self.klic}.xlsx"
        wb.save(cesta)
        soubory.insert(0, {"soubor": cesta.name, "list": None, "radku": None, "format": "xlsx"})
        return soubory


# ---------------------------------------------------------------------------
# Balíčky


def balicek_veletrhy(dnes: dt.date, rok_katalogu: str) -> tuple[Balicek, dict]:
    snimek = nacti_json(KOREN / "src" / "data" / "veletrhy-2027.json")
    potvrzene = [a for a in snimek["akce"] if a.get("terminPotvrzen") and a.get("start")]
    potvrzene.sort(key=lambda a: (a["start"], a.get("mesto") or "", a["id"]))
    kraje_nazvy = nacti_kraje(rok_katalogu)
    radky = []
    for a in potvrzene:
        konec = a.get("end") or a["start"]
        poznamka = []
        if a.get("terminPribligny"):
            poznamka.append("termín je přibližný")
        if a.get("zdrojJenAgregator"):
            poznamka.append("termín neověřený u pořadatele")
        if a.get("poznamkaTerminu"):
            poznamka.append(bez_relativniho_roku(a["poznamkaTerminu"], int(snimek["sezona"]) - 1))
        radky.append([
            nadpis_kraje(kraje_nazvy.get(a.get("krajKod"), a.get("krajKod") or "")),
            a.get("mesto") or ("online" if a.get("online") else ""),
            a.get("nazev"), a.get("datum"), a["start"], konec, a.get("cas") or "",
            a.get("misto") or "", a.get("poradatel") or "", a.get("url") or "",
            "; ".join(poznamka), "proběhla" if konec < dnes.isoformat() else "",
            a.get("overeno") or snimek.get("checkedAt") or "",
        ])
    po_krajich: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for a in potvrzene:
        k = nadpis_kraje(kraje_nazvy.get(a.get("krajKod"), a.get("krajKod") or ""))
        po_krajich[k][0] += 1
        if (a.get("end") or a["start"]) >= dnes.isoformat():
            po_krajich[k][1] += 1
    souhrn = [[k, v[0], v[1]] for k, v in sorted(po_krajich.items())]
    mesice = Counter(a["start"][:7] for a in potvrzene)
    b = Balicek(
        "veletrhy",
        f"Veletrhy a přehlídky středních škol, sezóna podzim {int(snimek['sezona']) - 1}",
        "Veletrh středních škol je akce, na které se na jednom místě představí víc středních škol najednou. "
        "Pořádá ho kraj, hospodářská komora, město nebo výstaviště.",
        [
            f"Stav k {dnes.isoformat()}. Obsahuje jen akce s potvrzeným termínem, tedy ověřeným na stránce pořadatele "
            f"({len(potvrzene)} akcí). Dalších {len(snimek['akce']) - len(potvrzene)} známých akcí na potvrzení termínu čeká.",
            "Přehled není úplný: obsahuje akce, o kterých víme. Kraj s málo akcemi není kraj s málo veletrhy.",
            "Kdo na akci vystavuje, žádný zdroj neuvádí; seznam vystavovatelů mají pořadatelé.",
            "Sloupec „proběhla“ je vyplněn u akcí, které k datu vytvoření balíčku skončily.",
            "Živý přehled: https://www.prijimackynaskolu.cz/veletrhy",
        ],
    )
    b.pridej("veletrhy-akce", "Akce", [
        "kraj", "mesto", "nazev", "termin", "od", "do", "cas", "misto", "poradatel", "odkaz",
        "poznamka_k_terminu", "proběhla", "overeno",
    ], radky)
    b.pridej("veletrhy-kraje", "Po krajích", [
        "kraj", "akci_s_potvrzenym_terminem_v_sezone", f"akci_od_{dnes.isoformat()}",
    ], souhrn)
    vysledek = {
        "akci_potvrzenych": len(potvrzene),
        "akci_cekajicich": len(snimek["akce"]) - len(potvrzene),
        "akci_od_dnes": sum(v[1] for v in po_krajich.values()),
        "kraju": len(po_krajich),
        "mesicu": dict(sorted(mesice.items())),
        "sezona": snimek["sezona"],
    }
    return b, vysledek


def katalog_nabidek(rok: str) -> dict:
    """Katalog nabídek 1. kola zobrazeného ročníku; název souboru nese rok z registru."""
    cesta = KOREN / "public" / f"applications_{rok}.json"
    katalog = nacti_json(cesta)
    if str(katalog["meta"]["rok"]) != rok:
        raise SystemExit(f"{cesta.name} nese rok {katalog['meta']['rok']}, registr {rok}")
    return katalog


def nacti_kraje(rok: str) -> dict[str, str]:
    """Kód NUTS → název kraje z katalogu nabídek (tentýž číselník jako CERMAT)."""
    data = katalog_nabidek(rok)["data"]
    return {r["kraj_kod"]: r["kraj"] for r in data if r.get("kraj_kod")}


def balicek_konzervatore(obec_kraj: dict[str, str], rok: str, kolo1: list[dict]) -> tuple[Balicek, dict]:
    index = nacti_json(KOREN / "data" / "msmt_rejstrik" / "nazvy-oboru.json")
    harmonogram = nacti_json(KOREN / "src" / "data" / "admissions-2027.json")
    skupina = next(g for g in harmonogram["groups"] if g["id"] == "konzervatore")
    obory_konz = ("82-44-P", "82-45-P", "82-46-P", "82-47-P")
    radky = []
    pocty = Counter()
    for redizo, obory in sorted(index["nabidky"].items(), key=lambda x: index["skoly"][x[0]][1] + index["skoly"][x[0]][0]):
        vlastni = [o for o in obory if o.startswith(obory_konz)]
        if not vlastni:
            continue
        nazev, obec = index["skoly"][redizo]
        ident = index["identifikace"].get(redizo, {})
        for o in vlastni:
            pocty[index["obory"].get(o, o)] += 1
        radky.append([
            nadpis_kraje(obec_kraj.get(obec, "")), obec, ident.get("uplny_nazev") or nazev, redizo,
            ident.get("adresa", ""),
            "; ".join(f"{index['obory'].get(o, o)} ({o})" for o in vlastni),
        ])
    terminy = [[e["title"], e["date"], e.get("note", "")] for e in skupina["events"]]
    # Výsledky 1. kola ze souhrnu CERMAT (obory konzervatoře nemají jednotnou zkoušku,
    # proto nejsou v katalogu webu). Všechny formy, forma je ve sloupci.
    vysledky = []
    celkem = Counter()
    for r in sorted((r for r in kolo1 if str(r["KKOV"]).startswith(obory_konz)),
                    key=lambda r: (r["KRAJ - NÁZEV"], r["OBEC"], r["NÁZEV ŠKOLY"], r["KKOV"], r.get("ZAMĚŘENÍ OBORU") or "")):
        radek = [cele(r[k]) for k in ("KAPACITA", "PŘIHLÁŠKY CELKEM", "PŘIJATÍ", "NEPŘIJATI - NEDOSTATEČNÁ KAPACITA",
                                       "NEPŘIJATI - NESPLNĚNÍ PODMÍNEK", "NEPŘIJATI - PŘIJAT NA VYŠŠÍ PRIORITU")]
        vysledky.append([
            nadpis_kraje(r["KRAJ - NÁZEV"]), r["OBEC"], r["NÁZEV ŠKOLY"], r["REDIZO"], r["OBOR - NÁZEV"], r["KKOV"],
            r.get("ZAMĚŘENÍ OBORU") or "", r["FORMA VZDĚLÁVÁNÍ"], cele(r["DÉLKA STUDIA"]), f"{r['ROČNÍK']}. třída", *radek,
        ])
        if denni_nezkracene_radek(r):
            for k, v in zip(("mist", "prihlasek", "prijatych", "neveslo", "nedosahlo", "vys"), radek):
                celkem[k] += v
    redizo_souhrn = {str(r["REDIZO"]) for r in kolo1 if str(r["KKOV"]).startswith(obory_konz)}
    redizo_rejstrik = {r[3] for r in radky}
    if redizo_souhrn != redizo_rejstrik:
        print(f"⚠️  Konzervatoře: rejstřík a souhrn {rok} se liší: {sorted(redizo_souhrn ^ redizo_rejstrik)}")
    obdobi = index["meta"]["obdobi"]
    b = Balicek(
        "konzervatore",
        "Konzervatoře: přihlášky do 30. listopadu",
        "Konzervatoře mají samostatné 1. kolo přijímacího řízení s podzimním termínem přihlášek. "
        "Ostatní střední školy včetně talentových oborů přijímají přihlášky až v únoru.",
        [
            f"Seznam konzervatoří je z rejstříku škol MŠMT ke dni {obdobi}: obsahuje školy, které mají zapsaný obor "
            "konzervatoře (hudba, zpěv, tanec, současný tanec, hudebně dramatické umění). Rejstřík říká, co škola smí učit, "
            "ne které obory pro nový ročník otevře a kolik přijme. Nabídku a kritéria zveřejní konzervatoře v DiPSy "
            f"v termínu {skupina['events'][0]['date']}.",
            f"Termíny jsou z harmonogramu MŠMT ({harmonogram['source']}), opsáno {harmonogram['checkedAt']}.",
            "Jednotná přijímací zkouška se u oborů konzervatoře nekoná; rozhoduje talentová zkouška.",
        ],
    )
    b.pridej("konzervatore-skoly", "Konzervatoře", ["kraj", "obec", "skola", "redizo", "adresa", "obory_z_rejstriku"], radky)
    b.pridej("konzervatore-terminy", "Termíny", ["udalost", "termin", "poznamka"], terminy)
    b.pridej(f"konzervatore-1-kolo-{rok}", f"1. kolo {rok}", [
        "kraj", "obec", "skola", "redizo", "obor", "kkov", "zamereni", "forma", "delka_studia", "hlasi_se_z",
        "mista", "prihlasky", "prijati", "nevesli_se_kvuli_kapacite", "nedosahli_pozadavku_skoly",
        "prijati_na_obor_vys_na_prihlasce",
    ], vysledky)
    b.o_datech.append(
        f"List „1. kolo {rok}“: souhrn CERMAT za 1. kolo {rok} (přihlášky v listopadu {int(rok) - 1}), všechny formy studia. "
        "Na obor Tanec se hlásí žáci 5. třídy, studium trvá osm let. Nedosáhli požadavku školy: typicky neprošli talentovou zkouškou. "
        "Přihlášky: jeden uchazeč podává víc přihlášek, proto se za různé obory nesčítají."
    )
    return b, {"skol": len(radky), "rejstrik_k": obdobi, "terminy": {e["id"]: e["date"] for e in skupina["events"]},
               "obory": dict(pocty), "kolo1_rok": rok, "kolo1_denni": dict(celkem), "shoda_se_souhrnem": redizo_souhrn == redizo_rejstrik}


def balicek_obory(rok: str, kolo1: list[dict], dk) -> tuple[Balicek, dict]:
    katalog = katalog_nabidek(rok)
    souhrny = nacti_json(KOREN / "public" / "souhrny_kolo1.json")["nabidky"]
    druhe = nacti_json(KOREN / "public" / "druhe_kolo.json")["roky"].get(rok, {})
    typy = {str(r.get("TYP ŠKOLY")): str(r.get("TYP ŠKOLY - NÁZEV")) for r in kolo1}
    # Okres po nabídkách včetně zaměření: jedna škola může mít obor ve více městech
    # (PORG, 79-41-K/81: Praha, Brno, Ostrava pod jedním REDIZO a KKOV).
    okresy = {dk.klic_webu(str(r["REDIZO"]), str(r["KKOV"]), str(r.get("ZAMĚŘENÍ OBORU") or "")): r.get("OKRES - NÁZEV")
              for r in kolo1 if str(r["POVINNOST JPZ"]).strip() == "1" and denni_nezkracene_radek(r)}
    obtiznost = {
        "kapacita_nerozhodovala": "místo bylo pro všechny, kdo splnili požadavky školy",
        "vetsina_uspela": "dostala se většina soutěžících uchazečů",
        "stredne_tezke": "středně těžké",
        "tezke": "těžké",
        "velmi_tezke": "velmi těžké",
    }
    stav2 = {"vypsano": "vypsáno", "nenaplneno_bez_2_kola": "nevypsáno, i když obor nebyl naplněn", "bez_2_kola": "nevypsáno"}
    # Oficiální nejnižší výsledek přijatých po nabídkách včetně zaměření (souhrn CERMAT,
    # sloupec ČJ+MA - % SKÓR - MIN (PŘIJATI) děleno dvěma), jen při aspoň deseti
    # přijatých s výsledkem zkoušky. Stejně jako u 2. kola; s vlastním výpočtem
    # z dat uchazečů se shoduje u 97 % oborů (docs/zdroje-dat.md, oddíl 2.1).
    minima = {}
    for r in kolo1:
        if str(r["POVINNOST JPZ"]).strip() != "1" or not denni_nezkracene_radek(r):
            continue
        m = cislo(r.get("ČJ+MA - % SKÓR - MIN (PŘIJATI)"))
        if m is not None and cele(r.get("ČJ+MA - KONALI (PŘIJATI)")) >= MIN_PRIJATYCH_PRO_MINIMUM:
            minima[dk.klic_webu(str(r["REDIZO"]), str(r["KKOV"]), str(r.get("ZAMĚŘENÍ OBORU") or ""))] = round(m / 2, 1)
    radky = []
    for n in sorted(katalog["data"], key=lambda x: (x["kraj"], x["obec"], x["nazev"], x["kkov"], x["zamereni"])):
        klic = dk.klic_webu(n["redizo"], n["kkov"], n["zamereni"])
        s = (souhrny.get(klic) or {}).get("roky", {}).get(rok, {})
        d = druhe.get(klic, {})
        minimum = minima.get(klic)
        radky.append([
            nadpis_kraje(n["kraj"]), okresy.get(klic) or "", n["obec"], n["nazev"], n["redizo"],
            n["obor"], n["kkov"], n["zamereni"], typy.get(n["typ"], n["typ"]), n.get("delka_studia"),
            n.get("kapacita"), s.get("prihlasky", n.get("prihlasky")),
            (s.get("prihlasky_priority") or [None])[0], s.get("index_poptavky"), s.get("tlak_prvnich_voleb"),
            s.get("prijati"), s.get("capacity_rejected"), s.get("conditions_not_met"), s.get("higher_priority"),
            s.get("podil_prijatych_ze_soutezicich"), obtiznost.get(s.get("zarazeni_obtiznosti"), ""),
            s.get("cj_ma_prijati"), minimum,
            stav2.get(d.get("stav"), ""), d.get("kapacita"), d.get("prihlasky"), d.get("prijati"),
        ])
    b = Balicek(
        f"obory-1-kolo-{rok}",
        f"Obory středních škol s jednotnou přijímací zkouškou: 1. kolo {rok}",
        f"Všech {len(radky)} nabídek denního nezkráceného studia s povinnou jednotnou přijímací zkouškou v 1. kole {rok}: "
        "poptávka, výsledek přijímání a zda škola vypsala 2. kolo.",
        [
            f"Zdroj: CERMAT, agregovaná data škol a oborů {rok}, platnost k {katalog['meta']['source']['valid_at']}; "
            f"2. kolo ze souhrnu 2. kola {rok}.",
            "Nepokrývá učební obory bez jednotné zkoušky, konzervatoře, dálkové a zkrácené studium.",
            f"Nabídku oborů pro nové přijímací řízení zveřejní školy spolu s kritérii přijetí. Obor z roku {rok} nemusí být vypsán znovu.",
            "Přihlášky: jeden uchazeč podává víc přihlášek, proto se přihlášky za různé obory nesčítají a přihlášky na místo konkurenci nadsazují.",
            "Tlak prvních voleb: kolik uchazečů chtělo obor jako 1. volbu na jedno místo.",
            "Soutěžící uchazeči jsou ti, kdo splnili požadavky školy a nedostali se na obor, který měli na přihlášce výš; "
            "podíl přijatých ze soutěžících říká, kolik z nich se dostalo. Není to šance konkrétního uchazeče.",
            "Body: součet bodů z češtiny a matematiky, každý test nejvýš 50 bodů. Průměr bodů přijatých neříká, s kolika body se dalo dostat. "
            f"Nejnižší výsledek přijatých je oficiální údaj CERMAT a uvádí se jen při aspoň {MIN_PRIJATYCH_PRO_MINIMUM} přijatých s výsledkem zkoušky; "
            "škola mohla vážit i jiná kritéria než test.",
            "Prázdná buňka znamená, že údaj zdroj nenese, ne nulu.",
        ],
    )
    b.pridej(f"obory-1-kolo-{rok}", "Obory", [
        "kraj", "okres", "obec", "skola", "redizo", "obor", "kkov", "zamereni", "typ_skoly", "delka_studia",
        "mista", "prihlasky", "prihlasky_jako_1_volba", "prihlasky_na_misto", "tlak_prvnich_voleb",
        "prijati", "nevesli_se_kvuli_kapacite", "nedosahli_pozadavku_skoly", "prijati_na_obor_vys_na_prihlasce",
        "podil_prijatych_ze_soutezicich", "obtiznost_prijeti", "prumer_bodu_prijatych", "nejnizsi_vysledek_prijatych",
        "2_kolo", "2_kolo_mista", "2_kolo_prihlasky", "2_kolo_prijati",
    ], radky)
    return b, {"nabidek": len(radky), "skol": len({r[4] for r in radky}), "prihlasek": katalog["meta"]["celkem_prihlasek"]}


def balicek_uchazeci(rok: str, uchazeci: list[dict], kolo1: list[dict]) -> tuple[Balicek, dict]:
    """Kam se uchazeči 1. kola dostali. Řádek zdroje je uchazeč.

    Populace: uchazeč s aspoň jednou přihláškou do denního nezkráceného studia mimo nástavbu
    (denni_nezkracene() ze scripts/slouceni_prihlasek.py). Kdo se hlásí jen na dálkové,
    kombinované či zkrácené studium nebo jen na nástavbu, není žák základní školy a počítá
    se zvlášť jako vyřazený. Přijetí se bere ze všech přihlášek uchazeče: přijetí na dálkové
    studium je také přijetí.

    Ročník, ze kterého se hlásí (5., 7. nebo 9. třída), se bere z přihlášek v populaci podle
    sloupce ROČNÍK souhrnu 1. kola. Kraj je kraj školy první takové přihlášky (bydliště
    uchazeče zdroj nenese). Pořadí volby je pořadí mezi všemi vyplněnými přihláškami,
    stejně jako v scripts/slouceni_prihlasek.py.
    """
    rocnik = {(str(r["REDIZO"]), str(r["KKOV"])): str(r["ROČNÍK"]) for r in kolo1}
    kraj = mapa_kraju(kolo1)
    tab: dict[tuple[str, str], Counter] = defaultdict(Counter)
    bez_rocniku = 0
    vyrazeno = Counter()
    for u in uchazeci:
        prihlasky = [k for k in range(1, 6) if u.get(f"ss{k}_redizo")]
        if not prihlasky:
            continue
        denni = [k for k in prihlasky if denni_nezkracene(u.get(f"ss{k}_forma"), u.get(f"ss{k}_zkraceno"))]
        v_populaci = [k for k in denni if not NASTAVBA.match(str(u[f"ss{k}_kkov"]))]
        if not v_populaci:
            vyrazeno["jen_nastavby" if denni else "jen_nedenni_nebo_zkracene"] += 1
            continue
        rocniky = {rocnik.get((str(u[f"ss{k}_redizo"]), str(u[f"ss{k}_kkov"]))) for k in v_populaci} - {None}
        if not rocniky:
            bez_rocniku += 1
            continue
        # Smíšené přihlášky (jednotky uchazečů) se řadí k nejvyššímu ročníku.
        roc = max(rocniky, key=int)
        prvni = v_populaci[0]
        kr = kraj.get((str(u[f"ss{prvni}_redizo"]), str(u[f"ss{prvni}_kkov"])), "")
        prijat_na = next((i for i, k in enumerate(prihlasky) if byl_prijat(u[f"ss{k}_prijat"], u[f"ss{k}_duvod_neprijeti"])), None)
        c = tab[(kr, roc)]
        c["uchazecu"] += 1
        c["konali_jpz"] += 1 if u.get("c_m_procentni_skor") not in (None, "") else 0
        if prijat_na is None:
            duvody = {u[f"ss{k}_duvod_neprijeti"] for k in prihlasky}
            c["nikam"] += 1
            c["nikam_s_jpz"] += 1 if u.get("c_m_procentni_skor") not in (None, "") else 0
            if duvody == {"pro_nedostacujici_kapacitu"}:
                c["nikam_jen_kapacita"] += 1
            elif duvody == {"pro_nesplneni_podminek"}:
                c["nikam_jen_pozadavek"] += 1
            elif duvody == {"pro_nedostacujici_kapacitu", "pro_nesplneni_podminek"}:
                c["nikam_obe"] += 1
            else:
                c["nikam_jine"] += 1
        else:
            c[f"volba{min(prijat_na + 1, 3)}"] += 1
    popis_rocniku = {"9": "9. třída (čtyřleté obory a obory bez maturity)", "7": "7. třída (šestiletá gymnázia)", "5": "5. třída (osmiletá gymnázia)"}
    hlavicka = [
        "kraj_skoly_1_volby", "rocnik", "uchazecu", "konali_jednotnou_zkousku", "prijati_na_1_volbu", "prijati_na_2_volbu",
        "prijati_na_3_a_dalsi_volbu", "neprijati_nikam", "podil_neprijatych_nikam",
        "z_toho_vsude_nevesli_kvuli_kapacite", "z_toho_vsude_nedosahli_pozadavku", "z_toho_obe_duvody", "z_toho_jiny_duvod",
    ]

    def radek(nazev: str, roc: str, c: Counter) -> list:
        return [nazev, popis_rocniku[roc], c["uchazecu"], c["konali_jpz"], c["volba1"], c["volba2"], c["volba3"],
                c["nikam"], podil(c["nikam"], c["uchazecu"]), c["nikam_jen_kapacita"], c["nikam_jen_pozadavek"],
                c["nikam_obe"], c["nikam_jine"]]

    celkem: dict[str, Counter] = defaultdict(Counter)
    radky = []
    for (kr, roc), c in sorted(tab.items(), key=lambda x: (x[0][0], -int(x[0][1]))):
        radky.append(radek(nadpis_kraje(kr), roc, c))
        celkem[roc].update(c)
    radky_cr = [radek("Česko celkem", roc, celkem[roc]) for roc in ("9", "7", "5")]
    b = Balicek(
        f"uchazeci-{rok}",
        f"Kam se dostali uchazeči v 1. a 2. kole {rok}",
        f"Uchazeči 1. kola {rok} podle toho, na kolikátou volbu z přihlášky byli přijati, nebo zda se nedostali nikam; "
        "uchazeči 2. kola a co víme o těch, kdo se nedostali ani ve 2. kole.",
        [
            f"Zdroj: CERMAT, data uchazečů 1. kola {rok}, předběžná verze: platné přihlášky ke dni 13. 5. {rok}. "
            "Jeden řádek zdroje je jeden uchazeč; zahrnuti jsou i uchazeči o obory bez jednotné zkoušky.",
            "Uchazeči, kteří se nedostali na víceleté gymnázium (5. a 7. třída), pokračují na základní škole. "
            "Proto se ročníky nesčítají a čísla za 9. třídu se nesmí míchat s víceletými gymnázii.",
            "Kraj je kraj školy, kterou měl uchazeč na přihlášce jako první. Bydliště uchazeče zdroj neuvádí.",
            "Volba: pořadí oboru na přihlášce. Pořadí na přihlášce šanci na přijetí nemění, škola řadí jen podle svých kritérií.",
            "Přijetí zahrnuje i uchazeče, kteří se přijetí později vzdali: data roku {rok} je nerozlišují.".format(rok=rok),
            "Nepřijati nikam: v 1. kole se nedostali na žádný obor z přihlášky. Mohli se hlásit do 2. kola (listy „2. kolo“).",
            "Důvody nepřijetí: nevešli se kvůli kapacitě (splnili požadavky školy, ale jiní měli lepší výsledek), "
            "nedosáhli požadavku školy (například minima bodů).",
        ],
    )
    b.pridej(f"uchazeci-1-kolo-{rok}-kraje", "Po krajích", hlavicka, radky)
    b.pridej(f"uchazeci-1-kolo-{rok}-cesko", "Česko", hlavicka, radky_cr)
    cr = {roc: dict(celkem[roc]) for roc in ("9", "7", "5")}
    kraje_9 = {nadpis_kraje(kr): dict(c) for (kr, roc), c in tab.items() if roc == "9"}
    b.o_datech.insert(1,
        f"Počítají se uchazeči s aspoň jednou přihláškou do denního nezkráceného studia mimo nástavbu. "
        f"Vyřazeno: {vyrazeno['jen_nedenni_nebo_zkracene']} uchazečů jen s přihláškami na dálkové, kombinované, "
        f"distanční, večerní nebo zkrácené studium a {vyrazeno['jen_nastavby']} jen s přihláškami na nástavbu; "
        "nejsou to žáci základní školy.")
    return b, {"rocniky": cr, "kraje_9": kraje_9, "bez_rocniku": bez_rocniku, "vyrazeno": dict(vyrazeno),
               "uchazecu": sum(c["uchazecu"] for c in celkem.values())}


UCEBNI = re.compile(r"-[CEHJ]/")  # kategorie bez maturity: učební obory a praktické školy


def druhe_kolo_uchazecu(b: Balicek, rok: str, uchazeci2: list[dict], kolo2: list[dict]) -> dict:
    """Doplní do balíčku uchazečů 2. kolo: kolik uchazečů se hlásilo a dostalo, kdo zůstal bez místa, volná místa.

    Zdroj je soubor uchazečů 2. kola (řádek = uchazeč, sloupec `rocnik` nese ročník přímo).
    Populace i počítání přijetí jsou stejné jako u 1. kola. Soubory 1. a 2. kola nemají
    společný identifikátor uchazeče, takže konkrétní dítě z 1. kola ve 2. kole nedohledáme;
    porovnávat jde jen počty.
    """
    kraj = mapa_kraju(kolo2)
    tab: dict[tuple[str, str], Counter] = defaultdict(Counter)
    for u in uchazeci2:
        prihlasky = [k for k in range(1, 6) if u.get(f"ss{k}_redizo")]
        v_populaci = [k for k in prihlasky if denni_nezkracene(u.get(f"ss{k}_forma"), u.get(f"ss{k}_zkraceno"))
                      and not NASTAVBA.match(str(u[f"ss{k}_kkov"]))]
        roc = str(u.get("rocnik") or "")
        if not v_populaci or roc not in ("5", "7", "9"):
            continue
        prvni = v_populaci[0]
        c = tab[(kraj.get((str(u[f"ss{prvni}_redizo"]), str(u[f"ss{prvni}_kkov"])), ""), roc)]
        c["uchazecu"] += 1
        c["s_vysledkem_jpz"] += 1 if u.get("c_m_procentni_skor") not in (None, "") else 0
        if any(byl_prijat(u[f"ss{k}_prijat"], u[f"ss{k}_duvod_neprijeti"]) for k in prihlasky):
            c["prijati"] += 1
            continue
        c["neprijati"] += 1
        duvody = {u[f"ss{k}_duvod_neprijeti"] for k in prihlasky}
        c["jen_kapacita" if duvody == {"pro_nedostacujici_kapacitu"} else "jen_pozadavek" if duvody == {"pro_nesplneni_podminek"}
          else "obe" if duvody == {"pro_nedostacujici_kapacitu", "pro_nesplneni_podminek"} else "jiny_duvod"] += 1
        ucebni = [bool(UCEBNI.search(str(u[f"ss{k}_kkov"]))) for k in v_populaci]
        c["jen_ucebni" if all(ucebni) else "jen_maturitni" if not any(ucebni) else "ucebni_i_maturitni"] += 1
        c["bez_vysledku_jpz"] += 1 if u.get("c_m_procentni_skor") in (None, "") else 0
        c["jedna_prihlaska"] += 1 if len(prihlasky) == 1 else 0

    # Volná místa po 2. kole: kapacita 2. kola bez přijatých, jen nabídky, které 2. kolo vypsaly.
    volno: dict[tuple[str, str, str], Counter] = defaultdict(Counter)
    for r in kolo2:
        if not denni_nezkracene_radek(r) or NASTAVBA.match(str(r["KKOV"])):
            continue
        zbylo = max(cele(r["KAPACITA"]) - cele(r["PŘIJATÍ"]), 0)
        c = volno[(nadpis_kraje(r["KRAJ - NÁZEV"]), str(r["ROČNÍK"]), "bez maturity" if UCEBNI.search(str(r["KKOV"])) else "s maturitou")]
        c["nabidek"] += 1
        c["volnych_mist"] += zbylo
        c["nabidek_s_volnym_mistem"] += 1 if zbylo else 0

    popis = {"9": "9. třída", "7": "7. třída (šestiletá gymnázia)", "5": "5. třída (osmiletá gymnázia)"}
    hl = ["kraj_skoly_1_volby", "rocnik", "uchazecu_2_kola", "prijati_ve_2_kole", "neprijati_ani_ve_2_kole",
          "podil_neprijatych", "z_toho_vsude_nevesli_kvuli_kapacite", "z_toho_vsude_nedosahli_pozadavku", "z_toho_obe_duvody",
          "z_toho_jiny_duvod", "z_toho_hlasili_se_jen_na_obory_bez_maturity", "z_toho_jen_na_maturitni", "z_toho_na_oboje",
          "z_toho_bez_vysledku_jednotne_zkousky", "z_toho_s_jedinou_prihlaskou"]

    def radek(nazev: str, roc: str, c: Counter) -> list:
        return [nazev, popis[roc], c["uchazecu"], c["prijati"], c["neprijati"], podil(c["neprijati"], c["uchazecu"]),
                c["jen_kapacita"], c["jen_pozadavek"], c["obe"], c["jiny_duvod"], c["jen_ucebni"], c["jen_maturitni"], c["ucebni_i_maturitni"],
                c["bez_vysledku_jpz"], c["jedna_prihlaska"]]

    celkem: dict[str, Counter] = defaultdict(Counter)
    radky = []
    for (kr, roc), c in sorted(tab.items(), key=lambda x: (x[0][0], -int(x[0][1]))):
        radky.append(radek(nadpis_kraje(kr), roc, c))
        celkem[roc].update(c)
    b.pridej(f"uchazeci-2-kolo-{rok}-kraje", "2. kolo po krajích", hl, radky)
    b.pridej(f"uchazeci-2-kolo-{rok}-cesko", "2. kolo Česko", hl, [radek("Česko celkem", r, celkem[r]) for r in ("9", "7", "5")])
    volno_celkem: dict[tuple[str, str], Counter] = defaultdict(Counter)
    for (_, roc, typ), c in volno.items():
        volno_celkem[(roc, typ)].update(c)
    b.pridej(f"volna-mista-po-2-kole-{rok}", "Volná místa po 2. kole",
             ["kraj", "rocnik", "obory", "nabidek_ve_2_kole", "nabidek_s_volnym_mistem", "volnych_mist"],
             [[k, popis.get(r, r), t, c["nabidek"], c["nabidek_s_volnym_mistem"], c["volnych_mist"]]
              for (k, r, t), c in sorted(volno.items())]
             + [["Česko celkem", popis.get(r, r), t, c["nabidek"], c["nabidek_s_volnym_mistem"], c["volnych_mist"]]
                for (r, t), c in sorted(volno_celkem.items())])
    b.o_datech.extend([
        f"Listy „2. kolo“: CERMAT, data uchazečů 2. kola {rok}, předběžná verze: platné přihlášky ke dni 23. 6. {rok}; "
        "stejná populace a stejné počítání přijetí jako u 1. kola.",
        "Soubory 1. a 2. kola nemají společný identifikátor uchazeče: konkrétní dítě z 1. kola ve 2. kole nedohledáme. "
        "Do 2. kola se navíc mohou přihlásit i ti, kdo v 1. kole přihlášku nepodali nebo se přijetí vzdali. "
        "Porovnávat jde jen počty, ne říct „z nepřijatých v 1. kole se ve 2. kole dostalo tolik“.",
        "Neprijati ani ve 2. kole: kam nastoupili, žádná zveřejněná data neříkají. Po 2. kole mohou školy vypisovat další kola "
        "na volná místa; data o nich CERMAT nezveřejňuje.",
        "Volná místa po 2. kole: kapacita 2. kola minus přijatí, jen u nabídek, které 2. kolo vypsaly (bez nástaveb). "
        "Neříká, zda škola místa nabídla v dalším kole, ani zda byla dostupná pro konkrétní dítě (kraj, obor, požadavky).",
    ])
    return {"rocniky": {r: dict(celkem[r]) for r in ("9", "7", "5")},
            "kraje_9": {nadpis_kraje(k): dict(c) for (k, r), c in tab.items() if r == "9"},
            "volno_9": {t: dict(c) for (r, t), c in volno_celkem.items() if r == "9"}}


def balicek_druhe_kolo(rok: str, kolo2: list[dict]) -> tuple[Balicek, dict]:
    radky = []
    souhrn_kraje: dict[tuple[str, str], Counter] = defaultdict(Counter)
    souhrn_typy: dict[str, Counter] = defaultdict(Counter)
    for r in sorted((r for r in kolo2 if denni_nezkracene_radek(r)),
                    key=lambda r: (r["KRAJ - NÁZEV"], r["OBEC"], r["NÁZEV ŠKOLY"], r["KKOV"], r.get("ZAMĚŘENÍ OBORU") or "")):
        jpz = str(r["POVINNOST JPZ"]).strip() == "1"
        kapacita, prihlasky, prijati = cele(r["KAPACITA"]), cele(r["PŘIHLÁŠKY CELKEM"]), cele(r["PŘIJATÍ"])
        neveslo = cele(r["NEPŘIJATI - NEDOSTATEČNÁ KAPACITA"])
        s_vysledkem = cele(r.get("ČJ+MA - KONALI (PŘIJATI)"))
        minimum = cislo(r.get("ČJ+MA - % SKÓR - MIN (PŘIJATI)"))
        minimum = round(minimum / 2, 1) if (jpz and minimum is not None and s_vysledkem >= MIN_PRIJATYCH_PRO_MINIMUM) else None
        radky.append([
            nadpis_kraje(r["KRAJ - NÁZEV"]), r.get("OKRES - NÁZEV") or "", r["OBEC"], r["NÁZEV ŠKOLY"], r["REDIZO"],
            r["OBOR - NÁZEV"], r["KKOV"], r.get("ZAMĚŘENÍ OBORU") or "", r["TYP ŠKOLY - NÁZEV"], cele(r["DÉLKA STUDIA"]),
            "ano" if jpz else "ne", kapacita, prihlasky, prijati, neveslo, cele(r["NEPŘIJATI - NESPLNĚNÍ PODMÍNEK"]),
            cele(r["NEPŘIJATI - PŘIJAT NA VYŠŠÍ PRIORITU"]), podil(prijati, kapacita, 2), minimum,
        ])
        for c in (souhrn_kraje[(nadpis_kraje(r["KRAJ - NÁZEV"]), "ano" if jpz else "ne")], souhrn_typy[r["TYP ŠKOLY - NÁZEV"]]):
            c["nabidek"] += 1
            c["mist"] += kapacita
            c["prihlasek"] += prihlasky
            c["prijatych"] += prijati
            c["neveslo"] += neveslo
            c["bez_prihlasky"] += 1 if prihlasky == 0 else 0
            c["s_nevesli"] += 1 if neveslo > 0 else 0

    def souhrn_radek(popis: list, c: Counter) -> list:
        return [*popis, c["nabidek"], c["mist"], c["prihlasek"], c["prijatych"], podil(c["prijatych"], c["mist"], 2),
                podil(c["prijatych"], c["prihlasek"], 2), c["neveslo"], c["s_nevesli"], c["bez_prihlasky"]]

    hl_s = ["nabidek", "mista", "prihlasky", "prijati", "naplnenost_mist", "prijatych_na_prihlasku",
            "nevesli_se_kvuli_kapacite", "nabidek_kde_se_nekdo_nevesel", "nabidek_bez_prihlasky"]
    kraje_radky = [souhrn_radek([k, j], c) for (k, j), c in sorted(souhrn_kraje.items())]
    typy_radky = [souhrn_radek([t], c) for t, c in sorted(souhrn_typy.items(), key=lambda x: -x[1]["mist"])]
    celkem = {j: Counter() for j in ("ano", "ne")}
    for (_, j), c in souhrn_kraje.items():
        celkem[j].update(c)
    b = Balicek(
        f"druhe-kolo-{rok}",
        f"Druhé kolo přijímacího řízení {rok}",
        "Když se obor v 1. kole nenaplní, škola může vypsat 2. kolo. Balíček obsahuje všechny nabídky 2. kola "
        f"{rok} v denním nezkráceném studiu, s jednotnou zkouškou i bez ní.",
        [
            f"Zdroj: CERMAT, agregovaná data škol a oborů, 2. kolo {rok}.",
            "Ve 2. kole se nová jednotná zkouška nepíše; škola bere výsledek z 1. kola nebo vlastní kritéria.",
            "Přihlášky 2. kola se nesčítají s přihláškami 1. kola a jeden uchazeč mohl podat víc přihlášek, "
            "proto přijatí na přihlášku nejsou podíl úspěšných uchazečů. Počty uchazečů 2. kola (lidí, ne přihlášek) "
            f"jsou v balíčku uchazeci-{rok}.",
            "Naplněnost míst: přijatí děleno místy 2. kola.",
            f"Nejnižší výsledek přijatých (body z češtiny a matematiky, 0–100) jen u oborů s jednotnou zkouškou a aspoň "
            f"{MIN_PRIJATYCH_PRO_MINIMUM} přijatými s výsledkem zkoušky. Neříká, s kolika body se dalo dostat.",
            "Vypsané 2. kolo neznamená, že ho škola vypíše znovu; obory, které se v 1. kole nenaplnily, ho vypsaly jen asi v polovině případů.",
        ],
    )
    b.pridej(f"druhe-kolo-{rok}-nabidky", "Nabídky", [
        "kraj", "okres", "obec", "skola", "redizo", "obor", "kkov", "zamereni", "typ_skoly", "delka_studia",
        "jednotna_zkouska", "mista", "prihlasky", "prijati", "nevesli_se_kvuli_kapacite", "nedosahli_pozadavku_skoly",
        "prijati_na_obor_vys_na_prihlasce", "naplnenost_mist", "nejnizsi_vysledek_prijatych",
    ], radky)
    b.pridej(f"druhe-kolo-{rok}-kraje", "Po krajích", ["kraj", "jednotna_zkouska", *hl_s], kraje_radky)
    b.pridej(f"druhe-kolo-{rok}-typy-skol", "Po typech škol", ["typ_skoly", *hl_s], typy_radky)
    kraje_jpz = {k: dict(c) for (k, j), c in souhrn_kraje.items() if j == "ano"}
    return b, {"jpz": dict(celkem["ano"]), "bez_jpz": dict(celkem["ne"]), "kraje_jpz": kraje_jpz,
               "typy": {t: dict(c) for t, c in souhrn_typy.items()}}


def skupiny_slozky(nazev: str) -> set[str]:
    return {k for k, vzor in SKUPINY_SLOZEK.items() if vzor.search(nazev or "")}


def balicek_kriteria(rok: str) -> tuple[Balicek, dict]:
    kriteria = nacti_json(KOREN / "public" / f"kriteria_prijeti_{rok}.json")
    katalog = {n["source_id"]: n for n in katalog_nabidek(rok)["data"]}
    vyhrada = VYHRADA_KRITERII.format(rok=rok)
    radky = []
    tab_typ: dict[str, Counter] = defaultdict(Counter)
    tab_kraj: dict[str, Counter] = defaultdict(Counter)
    nezarazene = Counter()
    for prepisy in kriteria["data"].values():
        for p in prepisy.get("prepisy", []):
            n = katalog.get(p.get("source_id"))
            if not n:
                continue
            skupiny: set[str] = set()
            for s in p.get("slozky", []):
                if NEJSOU_EXTRA_BODY.search((s.get("nazev") or "").strip()):
                    continue
                sk = skupiny_slozky(s.get("nazev", ""))
                if not sk and (s.get("max") or 0) > 0:
                    nezarazene[s.get("nazev", "")] += 1
                    sk = {"jine"}
                skupiny |= sk
            jen_jpz = p.get("rezim") == "pouze_jpz"
            radky.append([
                nadpis_kraje(n["kraj"]), n["obec"], n["nazev"], n["redizo"], n["obor"], n["kkov"], n["zamereni"], n["typ"],
                "ano" if jen_jpz else "ne", p.get("podil_jpz_pct"),
                *("ano" if k in skupiny else "" for k in SKUPINY_SLOZEK), "ano" if "jine" in skupiny else "",
                "; ".join(f"{s.get('nazev')} (max {s.get('max')})" for s in p.get("slozky", [])),
                "; ".join(p.get("minima") or []), "přepis ručně" if p.get("prepis") == "rucni" else "přepis strojově",
                "ano" if p.get("nalezy") else "", vyhrada,
            ])
            for c in (tab_typ[n["typ"]], tab_kraj[nadpis_kraje(n["kraj"])]):
                c["nabidek"] += 1
                c["jen_jpz"] += 1 if jen_jpz else 0
                for k in skupiny:
                    c[k] += 1
    radky.sort(key=lambda r: (r[0], r[1], r[2], r[5], r[6]))
    hl_skupiny = [f"{k}" for k in SKUPINY_SLOZEK]

    def souhrn(popis: str, c: Counter) -> list:
        return [popis, c["nabidek"], c["jen_jpz"], podil(c["jen_jpz"], c["nabidek"]), *(c[k] for k in SKUPINY_SLOZEK), c["jine"]]

    celkem = Counter()
    for c in tab_typ.values():
        celkem.update(c)
    hlavicka_s = ["skupina", "nabidek_s_prepisem", "bodovala_jen_jednotna_zkouska", "podil_jen_jednotna_zkouska",
                  *hl_skupiny, "jine_slozky"]
    b = Balicek(
        f"kriteria-{rok}",
        f"Co vedle jednotné přijímací zkoušky rozhodovalo: kritéria {rok}",
        f"Podle přepisu kritérií přijetí {rok} z DiPSy: u kterých oborů bodovala jen jednotná přijímací zkouška a kde "
        "škola přidávala extra body, tedy body za něco jiného, například za prospěch ze základní školy, "
        "školní přijímací zkoušku nebo pohovor.",
        [
            "VÝHRADA: " + vyhrada,
            f"Přepis vznikl strojově z PDF kritérií, která školy vložily do DiPSy pro 1. kolo {rok}; "
            f"pokrytí {kriteria['pokryti']['oboru_s_prepisem']} z {kriteria['pokryti']['oboru_s_pdf']} oborů s PDF.",
            "Skupiny složek (prospěch, školní zkouška, pohovor, talentová, praktická nebo sportovní zkouška a portfolio, soutěže a aktivity) "
            "přiřazuje generátor podle názvu složky v přepisu; jedna složka může patřit do víc skupin. "
            "Jiné složky jsou ty, které se podle názvu zařadit nepodařilo.",
            "Podíl přijímaček: kolik procent bodů celkového hodnocení tvořila jednotná zkouška podle přepisu. "
            "Neříká, jak moc další složky rozhodovaly: složka, kterou všichni dostanou plnou, pořadí nemění.",
            "Sloupce končící „_text_prepisu“ nesou slova z kritérií školy, jak je přepis převzal (například „minimální hranice přijetí“); "
            "nejsou to naše pojmy ani ověřené údaje.",
            "Jmenovité údaje o škole před zveřejněním ověřte v kritériích školy; výhrada je proto v každém řádku.",
            "Kritéria pro nové přijímací řízení zveřejní školy v DiPSy 15.–31. ledna.",
        ],
    )
    b.pridej(f"kriteria-{rok}-souhrn", "Souhrn", hlavicka_s,
             [souhrn("Česko celkem", celkem)]
             + [souhrn(f"typ {t}", c) for t, c in sorted(tab_typ.items())]
             + [souhrn(k, c) for k, c in sorted(tab_kraj.items())])
    b.pridej(f"kriteria-{rok}-obory", "Obory", [
        "kraj", "obec", "skola", "redizo", "obor", "kkov", "zamereni", "typ_skoly", "bodovala_jen_jednotna_zkouska",
        "podil_jednotne_zkousky_pct", *hl_skupiny, "jine_slozky", "slozky_text_prepisu", "minima_text_prepisu",
        "prepis", "mechanicka_kontrola_nasla_nesoulad", "vyhrada",
    ], radky)
    return b, {"nabidek": celkem["nabidek"], "jen_jpz": celkem["jen_jpz"], **{k: celkem[k] for k in SKUPINY_SLOZEK},
               "jine": celkem["jine"], "nezarazene_nazvy": nezarazene.most_common(30)}


# ---------------------------------------------------------------------------


def over_vstup(cesta: Path, ocekavany: str | None, popis: str) -> str:
    if not cesta.exists():
        raise SystemExit(f"Chybí {popis}: {cesta}")
    otisk = sha256(cesta)
    if ocekavany and otisk != ocekavany:
        raise SystemExit(f"{cesta.name}: otisk {otisk[:12]}… nesouhlasí s převzatým {ocekavany[:12]}… ({popis})")
    return otisk


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--vstupy", type=Path, required=True, help="adresář se zdrojovými XLSX CERMAT")
    ap.add_argument("--dnes", type=dt.date.fromisoformat, default=dt.date.today(), help="datum stavu balíčku")
    ap.add_argument("--vystup", type=Path, default=VYSTUP)
    args = ap.parse_args()

    registr = nacti_json(REGISTR)["sady"]
    rok = registr["cermat-vysledky"]["zobrazeno"]["obdobi"]
    rok_k2 = registr["cermat-kolo2-agregaty"]["zobrazeno"]["obdobi"]
    rok_u = registr["cermat-uchazeci-kolo1"]["zobrazeno"]["obdobi"]
    rok_kr = registr["dipsy-kriteria"]["zobrazeno"]["obdobi"]
    meta_vysledku = nacti_json(KOREN / "public" / "cermat_results_meta.json")["source"]

    kolo1_cesta = args.vstupy / f"PZ{rok}_kolo1_skolobory_vysledky.xlsx"
    kolo2_cesta = args.vstupy / f"PZ{rok_k2}_kolo2_skolobory_vysledky.xlsx"
    uch_cesta = args.vstupy / f"PZ{rok_u}_kolo1_uchazeci_prihlasky_vysledky.xlsx"
    rok_u2 = registr["cermat-uchazeci-kolo2"]["zobrazeno"]["obdobi"]
    uch2_cesta = args.vstupy / f"PZ{rok_u2}_kolo2_uchazeci_prihlasky_vysledky.xlsx"
    if rok_u2 != rok_k2:
        raise SystemExit(f"Uchazeči 2. kola ({rok_u2}) a souhrn 2. kola ({rok_k2}) musí být z téhož roku")
    otisky = {
        kolo1_cesta.name: over_vstup(kolo1_cesta, meta_vysledku.get("sha256"), "výsledky 1. kola"),
        kolo2_cesta.name: over_vstup(kolo2_cesta, registr["cermat-kolo2-agregaty"]["zobrazeno"].get("sha256"), "výsledky 2. kola"),
        # Data uchazečů nemají otisk v registru; ověřuje se počet řádků z poznámky registru a otisk se zapíše.
        uch_cesta.name: over_vstup(uch_cesta, None, "data uchazečů 1. kola"),
        uch2_cesta.name: over_vstup(uch2_cesta, registr["cermat-uchazeci-kolo2"]["zobrazeno"].get("sha256"), "data uchazečů 2. kola"),
    }

    kolo1 = nacti_xlsx(kolo1_cesta)
    kolo2 = nacti_xlsx(kolo2_cesta)
    uchazeci = nacti_xlsx(uch_cesta)
    uchazeci2 = nacti_xlsx(uch2_cesta)
    poznamka_u = registr["cermat-uchazeci-kolo1"]["zobrazeno"].get("z_dostupne", {}).get("poznamka", "")
    ocekavano_radku = int(re.sub(r"\D", "", poznamka_u.split("řádků")[0])) if "řádků" in poznamka_u else None
    if ocekavano_radku and len(uchazeci) != ocekavano_radku:
        raise SystemExit(f"Data uchazečů mají {len(uchazeci)} řádků, registr uvádí {ocekavano_radku}")

    nastaveni = nacti_json(NASTAVENI)
    dk = nacti_druhe_kolo_modul()
    obec_kraj = {str(r["OBEC"]): str(r["KRAJ - NÁZEV"]) for r in kolo1}

    balicky = []
    souhrn: dict = {}
    for nazev, (b, s) in [
        ("veletrhy", balicek_veletrhy(args.dnes, rok)),
        ("konzervatore", balicek_konzervatore(obec_kraj, rok, kolo1)),
        ("obory", balicek_obory(rok, kolo1, dk)),
        ("uchazeci", balicek_uchazeci(rok_u, uchazeci, kolo1)),
        ("druhe_kolo", balicek_druhe_kolo(rok_k2, kolo2)),
        ("kriteria", balicek_kriteria(rok_kr)),
    ]:
        if nazev == "uchazeci":
            # 2. kolo patří k 1. kolu téhož roku. V květnu a červnu je 1. kolo nového roku
            # převzaté dřív než 2. kolo; pak se 2. kolo vynechá, ne pod cizím rokem.
            s["kolo2"] = druhe_kolo_uchazecu(b, rok_u2, uchazeci2, kolo2) if rok_u2 == rok_u else None
            if s["kolo2"] is None:
                b.o_datech.append(f"Data uchazečů 2. kola {rok_u} zatím nejsou převzatá; balíček nese jen 1. kolo.")
        balicky.append((nazev, b))
        souhrn[nazev] = s

    args.vystup.mkdir(parents=True, exist_ok=True)
    for stary in args.vystup.glob("*"):
        if stary.suffix in (".csv", ".xlsx", ".zip", ".json"):
            stary.unlink()
    citace = nastaveni["citace_plna"] + f". Licence {nastaveni['licence']['nazev']}. {nastaveni['web']}/pro-novinare"
    katalog = []
    for nazev, b in balicky:
        soubory = b.zapis(args.vystup, citace)
        katalog.append({"klic": nazev, "soubor": b.klic, "nazev": b.nazev, "popis": b.popis, "soubory": soubory})

    zip_cesta = args.vystup / "prijimacky-na-skolu-data.zip"
    with zipfile.ZipFile(zip_cesta, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(args.vystup.glob("*")):
            if f.suffix in (".csv", ".xlsx"):
                z.write(f, f.name)
        z.writestr("CITACE.txt", citace + "\n")

    vystup = {
        "vytvoreno": args.dnes.isoformat(),
        "generator": "scripts/build-pro-novinare.py",
        "obdobi": {"vysledky": rok, "kolo2": rok_k2, "uchazeci": rok_u, "uchazeci_kolo2": rok_u2, "kriteria": rok_kr,
                   "veletrhy": souhrn["veletrhy"]["sezona"]},
        "zdroje": {"platnost_vysledku": meta_vysledku.get("valid_at"), "otisky": otisky},
        "balicky": katalog,
        "zip": zip_cesta.name,
        "skupiny_kriterii": NAZVY_SKUPIN,
        "cisla": {k: v for k, v in souhrn.items()},
    }
    nezarazene = vystup["cisla"]["kriteria"].pop("nezarazene_nazvy", [])
    (args.vystup / "souhrn.json").write_text(json.dumps(vystup, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    print(f"✅ {len(katalog)} balíčků do {args.vystup.relative_to(KOREN)}")
    u = souhrn["uchazeci"]["rocniky"]
    for roc in ("9", "7", "5"):
        c = u[roc]
        print(f"   {roc}. třída: {c['uchazecu']} uchazečů, nikam {c.get('nikam', 0)}")
    print("   2. kolo s JZ:", souhrn["druhe_kolo"]["jpz"])
    print("   kritéria:", {k: v for k, v in souhrn["kriteria"].items() if k != "nezarazene_nazvy"})
    print("   nezařazené složky kritérií (nejčastější):")
    for n, p in nezarazene[:40]:
        print(f"     {p}× {n}")


if __name__ == "__main__":
    main()
