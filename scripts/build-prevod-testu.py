#!/usr/bin/env python3
"""
Převodní tabulky výsledků cvičných testů CERMAT (TAU) na body zobrazeného roku.

Uchazeč si v aplikaci TAU udělá skutečný test z minulé jednotné zkoušky a zadá
body. Skript pro každý termín testu spočítá, na jakém místě mezi všemi
uchazeči, kteří ten test v ostrém termínu psali, by s tím výsledkem stál, a
najde body, které na stejném místě měl uchazeč v roce, za který web
zobrazuje pásma přijetí. Jde o ekvipercentilové převádění: předpokládá, že
pořadí se mezi ročníky zachovává, ne že jsou testy stejně těžké.

Zdroje:
  - položková data JPZ roku testu (data/JPZ{rok}_CJL4|MA4_polozkova_data.xlsx),
    jeden list na termín, řádek = uchazeč a test, `dt_body` 0–50,
  - data uchazečů cílového roku (data/PZ{cil}_kolo1_uchazeci_prihlasky_vysledky.xlsx),
    stejný soubor a stejná definice výsledku jako u pásem přijetí
    (`c_m_procentni_skor` / 2, lepší z pokusů).

Rok testu bere z registru (sada cermat-prevod-testu), cílový rok ze sady
cermat-uchazeci-kolo1. Výstup: public/prevod_testu_{rok}.json.

Definice a meze: slovník ukazatelů, *Převedený výsledek testu*.

    python3 scripts/build-prevod-testu.py
"""
from __future__ import annotations

import bisect
import json
import sys
from pathlib import Path

import openpyxl

KOREN = Path(__file__).resolve().parent.parent
REGISTR = KOREN / "public" / "stav_datovych_sad.json"

# Druhy testů podle délky studia: 4 = čtyřleté obory a nástavby (9. třída),
# 6 = šestiletá gymnázia (7. třída), 8 = osmiletá gymnázia (5. třída).
# Názvy souborů se mezi ročníky mění: 2024 CJL4/MA4, od 2025 C4/M4.
NAZVY_TESTU = {2024: {"4": ("CJL4", "MA4"), "6": ("CJL6", "MA6"), "8": ("CJL8", "MA8")}}


def druh_oboru(kkov: str) -> str:
    """Který test uchazeč o obor psal: víceletá gymnázia mají vlastní."""
    if kkov.endswith("K/81"):
        return "8"
    if kkov.endswith("K/61"):
        return "6"
    return "4"
# Listy jdou v pořadí termínů; název listu nese písmeno a popis.
TERMINY = [
    ("1-radny", "1. řádný termín", True),
    ("2-radny", "2. řádný termín", True),
    ("1-nahradni", "1. náhradní termín", False),
    ("2-nahradni", "2. náhradní termín", False),
]
# Pod tímto počtem řešitelů termínu je pořadí nespolehlivé (náhradní termíny 2024: 486 a 794).
MIN_RESITELU_SPOLEHLIVE = 5000


def obdobi(sada: str) -> int:
    registr = json.loads(REGISTR.read_text(encoding="utf-8"))
    return int(registr["sady"][sada]["zobrazeno"]["obdobi"])


def nacti_terminy(soubor: Path) -> list[dict[object, float]]:
    wb = openpyxl.load_workbook(soubor, read_only=True)
    listy = [ws for ws in wb.worksheets if "termín" in ws.title]
    if len(listy) != 4:
        sys.exit(f"{soubor.name}: čekám 4 listy s termíny, našel jsem {len(listy)}.")
    vysledky = []
    for ws in listy:
        it = ws.iter_rows(values_only=True)
        ix = {n: i for i, n in enumerate(next(it))}
        # Identifikátor uchazeče: v češtině čtyřletých `id_ss`, jinde `id`.
        sloupec_id = "id_ss" if "id_ss" in ix else "id"
        body = {}
        for r in it:
            if r[ix["dt_body"]] is not None:
                body[r[ix[sloupec_id]]] = float(r[ix["dt_body"]])
        vysledky.append(body)
    return vysledky


def nacti_cil(soubor: Path) -> dict[str, tuple[list[float], list[float], list[float]]]:
    """Rozdělení cílového roku zvlášť podle druhu testu.

    Uchazeč o víceleté gymnázium psal jiný test než uchazeč o čtyřletý obor;
    smíchané rozdělení by převod zkreslilo (2026: 25 400 z 119 000 uchazečů
    s výsledkem jsou uchazeči o víceletá gymnázia). Druh se určí z oborů na
    přihlášce.
    """
    wb = openpyxl.load_workbook(soubor, read_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    ix = {n: i for i, n in enumerate(next(it))}
    vystup = {d: ([], [], []) for d in ("4", "6", "8")}
    for r in it:
        obory = [r[ix[f"ss{k}_kkov"]] for k in range(1, 6) if r[ix[f"ss{k}_kkov"]]]
        druhy = {druh_oboru(o) for o in obory}
        if len(druhy) != 1:
            continue  # bez oboru nebo nejednoznačně: do žádného rozdělení
        celkem, cj, ma = vystup[druhy.pop()]
        if r[ix["c_m_procentni_skor"]] is not None:
            celkem.append(float(r[ix["c_m_procentni_skor"]]) / 2)
        if r[ix["c_procentni_skor"]] is not None:
            cj.append(float(r[ix["c_procentni_skor"]]) / 2)
        if r[ix["m_procentni_skor"]] is not None:
            ma.append(float(r[ix["m_procentni_skor"]]) / 2)
    return {d: (sorted(a), sorted(b), sorted(c)) for d, (a, b, c) in vystup.items()}


def poradi(serazene: list[float], body: float) -> float:
    """Percentil se středním pořadím: polovina shodných se počítá pod."""
    pod = bisect.bisect_left(serazene, body)
    shodnych = bisect.bisect_right(serazene, body) - pod
    return (pod + shodnych / 2) / len(serazene)


def kvantil(serazene: list[float], p: float) -> float:
    """Hodnota na daném podílu pořadí, s lineární interpolací."""
    poloha = min(max(p, 0.0), 1.0) * (len(serazene) - 1)
    dole = int(poloha)
    nahore = min(dole + 1, len(serazene) - 1)
    return serazene[dole] + (serazene[nahore] - serazene[dole]) * (poloha - dole)


def tabulka(zdroj: list[float], cil: list[float], maximum: int) -> dict:
    zdroj = sorted(zdroj)
    percentily, body = [], []
    for b in range(maximum + 1):
        p = poradi(zdroj, b)
        percentily.append(round(p * 100, 1))
        body.append(round(kvantil(cil, p), 1))
    return {"percentil": percentily, "body_cil": body}


def main() -> None:
    rok = obdobi("cermat-prevod-testu")
    cil_rok = obdobi("cermat-uchazeci-kolo1")
    if rok not in NAZVY_TESTU:
        sys.exit(f"Pro rok {rok} neznám názvy testů; doplňte NAZVY_TESTU.")
    cile = nacti_cil(KOREN / "data" / f"PZ{cil_rok}_kolo1_uchazeci_prihlasky_vysledky.xlsx")

    druhy = {}
    for druh, (cj_nazev, ma_nazev) in NAZVY_TESTU[rok].items():
        cj_terminy = nacti_terminy(KOREN / "data" / f"JPZ{rok}_{cj_nazev}_polozkova_data.xlsx")
        ma_terminy = nacti_terminy(KOREN / "data" / f"JPZ{rok}_{ma_nazev}_polozkova_data.xlsx")
        cil_celkem, cil_cj, cil_ma = cile[druh]
        terminy = []
        for (klic, nazev, radny), cj, ma in zip(TERMINY, cj_terminy, ma_terminy):
            # Součet jen u těch, kdo v tomtéž termínu psali oba testy.
            spolecni = cj.keys() & ma.keys()
            soucty = [cj[i] + ma[i] for i in spolecni]
            terminy.append({
                "klic": klic,
                "nazev": nazev,
                "radny": radny,
                "resitelu": len(spolecni),
                "spolehlive": len(spolecni) >= MIN_RESITELU_SPOLEHLIVE,
                "celkem": tabulka(soucty, cil_celkem, 100),
                "cj": tabulka(list(cj.values()), cil_cj, 50),
                "ma": tabulka(list(ma.values()), cil_ma, 50),
            })
            print(f"druh {druh}, {nazev}: {len(spolecni)} řešitelů; 50 bodů → "
                  f"{terminy[-1]['celkem']['body_cil'][50]} bodů roku {cil_rok} (cíl {len(cil_celkem)} uchazečů)")
        druhy[druh] = {
            "zdroj_testu": f"JPZ{rok}_{cj_nazev}_polozkova_data.xlsx, JPZ{rok}_{ma_nazev}_polozkova_data.xlsx",
            "cil_uchazecu": len(cil_celkem),
            "terminy": terminy,
        }

    vystup = {
        "rok_testu": rok,
        "rok_cile": cil_rok,
        "zdroj_cile": f"PZ{cil_rok}_kolo1_uchazeci_prihlasky_vysledky.xlsx",
        "metoda": "ekvipercentilové převádění: pořadí mezi řešiteli testu (střední pořadí) → body se stejným pořadím mezi uchazeči cílového roku, kteří psali tentýž druh testu",
        "druhy_testu": {"4": "čtyřleté obory a nástavby (9. třída)", "6": "šestiletá gymnázia (7. třída)", "8": "osmiletá gymnázia (5. třída)"},
        "druhy": druhy,
    }
    cesta = KOREN / "public" / f"prevod_testu_{rok}.json"
    cesta.write_text(json.dumps(vystup, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"Zapsáno {cesta.relative_to(KOREN)} ({cesta.stat().st_size // 1024} kB)")


if __name__ == "__main__":
    main()
