"""Prezentace AI First Company pro Vibecoding Talks (~10 min, 16:9)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

import os

# Podklady leží vedle skriptu, ať jde prezentace přegenerovat odkudkoli.
ZDE = os.path.dirname(os.path.abspath(__file__))
S = os.path.join(ZDE, 'podklady')

TMAVA = RGBColor(0x16, 0x32, 0x5C)
MODRA = RGBColor(0x00, 0x74, 0xE4)
SEDA = RGBColor(0x5B, 0x68, 0x77)
SVETLA = RGBColor(0xF4, 0xF7, 0xFB)
BILA = RGBColor(0xFF, 0xFF, 0xFF)
CERVENA = RGBColor(0xC0, 0x39, 0x2B)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
PRAZDNY = prs.slide_layouts[6]


def slide(pozadi=BILA):
    s = prs.slides.add_slide(PRAZDNY)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = pozadi
    return s


def text(s, x, y, w, h, obsah, velikost=20, tucne=False, barva=TMAVA,
         zarovnani=PP_ALIGN.LEFT, proklad=1.25, kurziva=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, radek in enumerate(obsah.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = zarovnani
        p.line_spacing = proklad
        r = p.add_run()
        r.text = radek
        r.font.size = Pt(velikost)
        r.font.bold = tucne
        r.font.italic = kurziva
        r.font.color.rgb = barva
        r.font.name = 'Calibri'
    return tb


def karta(s, x, y, w, h, vypln=SVETLA, obrys=None):
    from pptx.enum.shapes import MSO_SHAPE
    tv = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    tv.fill.solid()
    tv.fill.fore_color.rgb = vypln
    if obrys:
        tv.line.color.rgb = obrys
        tv.line.width = Pt(1.5)
    else:
        tv.line.fill.background()
    tv.shadow.inherit = False
    tv.text_frame.text = ''
    return tv


def obrazek_vyska(s, cesta, x, y, vyska_in):
    """Vloží obrázek o dané výšce, šířku dopočítá z poměru stran."""
    im = Image.open(cesta)
    pomer = im.width / im.height
    return s.shapes.add_picture(cesta, Inches(x), Inches(y), height=Inches(vyska_in),
                                width=Inches(vyska_in * pomer))


def poznamka(s, txt):
    s.notes_slide.notes_text_frame.text = txt


def patka(s, cislo):
    text(s, 11.9, 6.95, 1.2, 0.4, str(cislo), 11, barva=SEDA, zarovnani=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------- 1. Titulka
s = slide(TMAVA)
text(s, 1.0, 2.3, 11.3, 1.2, 'AI First Company', 54, True, BILA)
text(s, 1.0, 3.5, 11.3, 1.0, 'Přijímačky na školu jako provozní experiment', 26, False,
     RGBColor(0xC9, 0xD4, 0xE1))
text(s, 1.0, 4.9, 11.3, 0.6, 'Patrick Zandl  ·  Vibecoding Talks', 16, False,
     RGBColor(0x8C, 0xA3, 0xBD))
poznamka(s, 'Nebudu prodávat, že to funguje. Ukážu, kde to selhalo — a proč právě to je '
            'důkaz, že model drží. 10 minut, rezerva do 15.')

# ---------------------------------------------------------------- 2. Hook
s = slide()
text(s, 1.0, 1.5, 11.3, 1.6, 'Minulý týden jsem řízení firmy\npředal AI ředitelce.', 40, True)
text(s, 1.0, 3.6, 11.3, 0.8, 'Hype, nebo provozovatelný model?', 28, False, MODRA)
karta(s, 1.0, 4.8, 11.3, 1.5)
text(s, 1.4, 5.05, 10.5, 1.1,
     'Co si odnesete: provozní rámec „agent navrhuje, člověk schvaluje následky“,\n'
     'tři konkrétní průšvihy — a co je zachytilo.', 18, False, SEDA)
patka(s, 2)
poznamka(s, 'Otázka do sálu. Nechat vydechnout. Pak rovnou k zadání.')

# ---------------------------------------------------------------- 3. Zadání
s = slide()
text(s, 1.0, 0.7, 11.3, 0.7, 'Zadání, ne roadmapa', 36, True)
karta(s, 1.0, 2.0, 5.4, 2.6, SVETLA)
text(s, 1.4, 2.35, 4.7, 2.0, 'Cíl', 16, True, MODRA)
text(s, 1.4, 2.85, 4.7, 1.6,
     'Poskytovat co nejlepší informace\no přijímacím řízení na střední školy.', 19)
karta(s, 6.9, 2.0, 5.4, 2.6, RGBColor(0xFD, 0xF2, 0xF0), CERVENA)
text(s, 7.3, 2.35, 4.7, 2.0, 'Mez', 16, True, CERVENA)
text(s, 7.3, 2.85, 4.7, 1.6,
     'Bez pokusu o průnik na cizí servery.\n(doplněno po rozhovoru Slížek / Koubský)', 19)
text(s, 1.0, 5.2, 11.3, 1.0,
     'Žádný 40slidový plán. Agent dostane cíl a meze, ne todo list.', 24, True, TMAVA)
patka(s, 3)
poznamka(s, 'Klíč: meze jsou součástí zadání, ne dodatečná kontrola.')

# ---------------------------------------------------------------- 4. Organigram
s = slide()
text(s, 0.7, 0.45, 12.0, 0.7, 'Dělba moci, ne delegace', 36, True)
text(s, 0.7, 1.15, 12.0, 0.5, 'Čtyři role, tři modely, jeden člověk', 18, False, SEDA)

sloupce = [
    ('Chief of Staff', 'orchestrace,\nprovozní agenda', RGBColor(0x6B, 0x4E, 0xA8)),
    ('Eduarda', 'ředitelka — návrhy,\nkomunikace, inbox', MODRA),
    ('Claude Code', 'realizace\npodle zadání', RGBColor(0x1E, 0x7A, 0x5E)),
    ('Codex', 'oponentura,\nnálezy', RGBColor(0xC2, 0x6A, 0x1A)),
]
x = 0.7
for nazev, popis, barva in sloupce:
    karta(s, x, 2.0, 2.75, 2.1, SVETLA, barva)
    if nazev == 'Eduarda':
        # Portrét patří vedle ostatních agentů, ne samostatně jako fotka
        # zaměstnankyně — v talku o AI First je ten rozdíl podstatný.
        s.shapes.add_picture(f'{S}/eduarda-portret.png', Inches(x + 1.92),
                             Inches(2.12), height=Inches(0.72), width=Inches(0.72))
    text(s, x + 0.25, 2.25, 2.3, 0.5, nazev, 17, True, barva)
    text(s, x + 0.25, 2.8, 2.3, 1.2, popis, 13, False, SEDA, proklad=1.15)
    if x < 9.5:
        text(s, x + 2.72, 2.8, 0.5, 0.5, '▶', 14, False, RGBColor(0xC9, 0xD4, 0xE1))
    x += 3.0

karta(s, 0.7, 4.45, 11.9, 0.95, TMAVA)
text(s, 1.1, 4.62, 11.1, 0.7, 'Patrick  ·  schvaluje realizaci — peníze, právo, reputace',
     22, True, BILA)
text(s, 0.7, 5.75, 11.9, 0.9,
     '„Kdo kód píše, ten ho nerecenzuje. Kdo zadává, ten neschvaluje.“',
     26, True, MODRA, PP_ALIGN.CENTER)
text(s, 0.7, 6.55, 11.9, 0.5, 'Eduarda běží na Grok Botu — tam má skills, napojení na svět '
     'i knowledge base.', 14, False, SEDA, PP_ALIGN.CENTER)
patka(s, 4)
poznamka(s, 'JÁDRO PREZENTACE. Každá vrstva má jiný zájem: co se má stát / jak to postavit / '
            'co je na tom špatně / co to stojí a čím to hrozí. Ověřit před talkem, jestli je '
            'Chief of Staff samostatný agent nebo řídicí vrstva.')

# ---------------------------------------------------------------- 5. Eduarda v akci
s = slide()
text(s, 0.7, 0.45, 7.3, 0.7, 'Odpověděla. Předala. Zeptala se.', 36, True)
text(s, 0.7, 1.2, 7.3, 0.5, 'Jedna kontrola inboxu, tři různá chování', 18, False, SEDA)

bloky = [
    ('Odpověděla sama', 'Tip na kalendář dnů otevřených dveří — zodpovězeno z eda@,\n'
     'založen produktový nápad.', RGBColor(0x1E, 0x7A, 0x5E)),
    ('Předala výš', 'DMARC reporty → Chief of Staff.\nNení to její agenda.',
     RGBColor(0x6B, 0x4E, 0xA8)),
    ('Zeptala se', '„Jestli to měl být ostrý test podpory, napiš\na odpovím jako zákazníkovi.“',
     MODRA),
]
y = 2.0
for nazev, popis, barva in bloky:
    karta(s, 0.7, y, 6.9, 1.45, SVETLA, barva)
    text(s, 1.05, y + 0.15, 6.3, 0.45, nazev, 17, True, barva)
    text(s, 1.05, y + 0.65, 6.3, 0.8, popis, 13, False, SEDA, proklad=1.15)
    y += 1.6

text(s, 0.7, 6.85, 7.3, 0.5, 'Autonomie roste s vratností, ne s důvěrou.', 17, True, TMAVA)
obrazek_vyska(s, f'{S}/eduarda-anonym.png', 8.5, 0.5, 6.6)
patka(s, 5)
poznamka(s, 'Nejsilnější slide. Nechat publikum chvíli číst. Zvlášť vypíchnout: „Když nic '
            'nového, v chatu zůstávám ticho.“ Agent, který mlčí, když nemá co říct.')

# ---------------------------------------------------------------- 6. Týden v číslech
s = slide()
text(s, 0.7, 0.5, 12.0, 0.7, 'Týden v provozu', 36, True)
cisla = [
    ('20 / 20', 'pozvánek doručeno\n0 odmítnutí'),
    ('368', 'automatických\ntestů'),
    ('23', 'pořadatelů akcí\nv databázi'),
    ('150 Kč', 'denní limit\nna reklamu'),
]
x = 0.7
for velke, popis in cisla:
    karta(s, x, 1.7, 2.75, 2.2, SVETLA)
    text(s, x + 0.2, 1.95, 2.35, 0.9, velke, 40, True, MODRA, PP_ALIGN.CENTER)
    text(s, x + 0.2, 2.95, 2.35, 0.9, popis, 13, False, SEDA, PP_ALIGN.CENTER, proklad=1.15)
    x += 3.0

text(s, 0.7, 4.2, 12.0, 2.0,
     'Portál škol — profily, které si školy samy upravují; publikace ihned, moderace zpětně\n'
     'Sklízení dat z webů škol přes RSS a Exa\n'
     'Odhad nákladů, kdyby totéž zadalo MŠMT:  ~5 mil. Kč', 20, False, TMAVA, proklad=1.6)
patka(s, 6)
poznamka(s, 'Čísla jsou ověřená, ne odhadnutá. 368 testů = skutečný stav repozitáře. '
            '20/20 doručeno = potvrzeno Resendem.')

# ---------------------------------------------------------------- 7. Marketing
s = slide()
text(s, 0.7, 0.45, 6.6, 0.7, 'Výstup, ne slib', 36, True)
text(s, 0.7, 1.25, 6.6, 2.4,
     'Zadání znělo: seber databázi pořadatelů\nveletrhů středních škol.\n\n'
     'Výstup: 23 organizací, u 12 ověřený e-mail,\nExcel i CSV, seřazené podle\n'
     'nejbližšího deadlinu.', 20, False, TMAVA, proklad=1.35)
karta(s, 0.7, 4.2, 6.6, 1.5, SVETLA, MODRA)
text(s, 1.05, 4.45, 6.0, 1.1,
     'Agent nedodal rešerši k dalšímu zpracování.\nDodal seznam, ze kterého se dá rovnou psát.',
     16, False, SEDA, proklad=1.25)
text(s, 0.7, 6.3, 6.6, 0.6, 'Adresy na slidu zakryté — nemám souhlas jejich majitelů.',
     13, False, SEDA, kurziva=True)
obrazek_vyska(s, f'{S}/marketing-anonym.png', 8.3, 0.5, 6.6)
patka(s, 7)
poznamka(s, 'DOPLNIT: jak dlouho to trvalo — to číslo neznám, nevymýšlet si ho na pódiu. '
            'Zmínit, proč jsou adresy zakryté — publikum to ocení a je to malá ukázka téhož '
            'principu: co má následky, řeší člověk.')

# ---------------------------------------------------------------- 8. Kde to drhlo
s = slide()
text(s, 0.7, 0.45, 12.0, 0.7, 'Kde to drhlo', 36, True)
text(s, 0.7, 1.2, 12.0, 0.5, 'Tři nálezy z jednoho pull requestu', 18, False, SEDA)

chyby = [
    ('Vymyšlený technický fakt',
     'AI tvrdila, že tečka za kódem rozbije přihlášení. Nerozbije — normalizace ji zahodí.\n'
     'Ta nepravda se dostala do dokumentu, který byl zdrojem pravdy pro rozesílku.'),
    ('Vývojářská poznámka mířila do e-mailu',
     'HTML komentář „ředitel by si kód zkopíroval špatně“ by dorazil dvaceti ředitelům —\n'
     've studeném e-mailu, jehož celý smysl je nevypadat jako podvod.'),
    ('Jednorázový kód se dal spálit',
     '253 z 1362 škol. Kód je jednorázový; druhý už škola nedostane.'),
]
y = 1.9
for nazev, popis in chyby:
    karta(s, 0.7, y, 11.9, 1.45, RGBColor(0xFD, 0xF2, 0xF0), CERVENA)
    text(s, 1.05, y + 0.15, 11.2, 0.45, nazev, 19, True, CERVENA)
    text(s, 1.05, y + 0.68, 11.2, 0.8, popis, 14, False, SEDA, proklad=1.2)
    y += 1.6

text(s, 0.7, 6.75, 11.9, 0.6,
     'Nic z toho nezachytila opatrnost. Zachytilo to další kolo.', 24, True, TMAVA,
     PP_ALIGN.CENTER)
patka(s, 8)
poznamka(s, 'Tohle je ta část, kvůli které talk není sales pitch. Nešetřit se.')

# ---------------------------------------------------------------- 9. Co to chytilo
s = slide()
text(s, 0.7, 0.5, 12.0, 0.7, 'Co to zachytilo', 36, True)
karta(s, 0.7, 1.6, 5.8, 2.3, SVETLA, MODRA)
text(s, 1.05, 1.85, 5.2, 0.5, '6 kol', 40, True, MODRA)
text(s, 1.05, 2.6, 5.2, 1.2, 'code review nad jedním\npull requestem', 17, False, SEDA)
karta(s, 6.8, 1.6, 5.8, 2.3, SVETLA, CERVENA)
text(s, 7.15, 1.85, 5.2, 0.5, '9 testů', 40, True, CERVENA)
text(s, 7.15, 2.6, 5.2, 1.2, 'netestovalo vůbec nic —\nodhalila to mutace', 17, False, SEDA)

text(s, 0.7, 4.3, 11.9, 1.8,
     'Jeden z nich chybu dokonce zamykal jako správné chování.\n\n'
     'Model, který kód napsal, je na něj slepý. Tři z těch nálezů byly věci,\n'
     'které si byl jistý — sám by je neodhalil.', 21, False, TMAVA, proklad=1.4)
text(s, 0.7, 6.5, 11.9, 0.7,
     'Nestav kontrolu na tom, že agent bude opatrný. Postav ji na tom, že se dívá někdo jiný.',
     20, True, MODRA, PP_ALIGN.CENTER)
patka(s, 9)
poznamka(s, 'Mutace = rozbij kód, který má test hlídat. Když projde, netestuje nic. '
            'Nejlevnější pojistka, jakou znám.')

# ---------------------------------------------------------------- 10. Hranice
s = slide()
text(s, 0.7, 0.5, 12.0, 0.7, 'Kde končí autonomie', 36, True)
karta(s, 0.7, 1.7, 5.8, 3.6, RGBColor(0xEF, 0xF7, 0xF1), RGBColor(0x1E, 0x7A, 0x5E))
text(s, 1.1, 1.95, 5.1, 0.5, 'Agent sám', 20, True, RGBColor(0x1E, 0x7A, 0x5E))
text(s, 1.1, 2.55, 5.1, 2.6,
     'Kód a refaktoring\nOdpovědi v podpoře\nRešerše a databáze\nKampaně v rámci limitu\n'
     'Návrhy produktu', 17, False, SEDA, proklad=1.55)
karta(s, 6.8, 1.7, 5.8, 3.6, RGBColor(0xFD, 0xF2, 0xF0), CERVENA)
text(s, 7.2, 1.95, 5.1, 0.5, 'Vždy člověk', 20, True, CERVENA)
text(s, 7.2, 2.55, 5.1, 2.6,
     'Platby\nSmlouvy\nÚřady\nNevratné rozesílky\nCokoli s dopadem na reputaci',
     17, False, SEDA, proklad=1.55)
text(s, 0.7, 5.6, 11.9, 1.2,
     'Dělicí čára není „důležité / nedůležité“.\nJe to „dá se to vzít zpět?“',
     24, True, TMAVA, PP_ALIGN.CENTER, proklad=1.3)
patka(s, 10)
poznamka(s, 'Poctivě přiznat: za ten týden jsem do toho mluvil víc, než zní „já jen schvaluju '
            'platby“. Autonomie zatím není u strategie, ale u realizace.')

# ---------------------------------------------------------------- 11. Takeaway
s = slide(TMAVA)
text(s, 1.0, 1.1, 11.3, 0.8, 'Co zkusit zítra', 36, True, BILA)
kroky = [
    ('1', 'Dej agentovi cíl a meze, ne seznam úkolů.'),
    ('2', 'Rozděl role: kdo píše, ten nerecenzuje.'),
    ('3', 'U každého testu rozbij kód, který má hlídat.\nKdyž projde, netestuje nic.'),
]
y = 2.3
for cislo, txt in kroky:
    text(s, 1.0, y, 0.7, 0.7, cislo, 30, True, MODRA)
    text(s, 1.8, y + 0.05, 10.3, 1.0, txt, 21, False, RGBColor(0xE3, 0xE9, 0xF1), proklad=1.25)
    y += 1.35

text(s, 1.0, 6.5, 11.3, 0.6, 'eda@prijimackynaskolu.cz  ·  thread na X', 17, False,
     RGBColor(0x8C, 0xA3, 0xBD))
poznamka(s, 'Zakončit pozvánkou: business návrhy Eduardě na eda@. Bacha, taky u toho mailu '
            'nesedí pořád, má práci.')

prs.save(os.path.join(ZDE, 'ai-first-company.pptx'))
print(f'uloženo: {len(prs.slides.__iter__.__self__._sldIdLst)} slidů')
