"""Testy importu souhrnů CERMATu (scripts/import_cermat_results.py, refresh_cermat_data.py).

Psáno v unittestu, ne v pytestu: projekt jinde používá unittest a CI instaluje
jen openpyxl (.github/workflows/datova-linka.yml), takže pytestová verze tohoto
souboru se nikde nespouštěla a chyběla i v `python3 -m unittest discover -s tests`.
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent.parent / 'scripts'))
from import_cermat_results import load_flat_xlsx, extract_current_year, compute_ranks, make_key, is_valid_flat
from refresh_cermat_data import build_applications


def record(**overrides):
    return {
        'ROK': 2026, 'KOLO': 1, 'REDIZO': 123, 'KKOV': '79-41-K/41', 'ID_SOF': '123-x',
        'FORMA VZDĚLÁVÁNÍ': 'denní', 'ZKRÁCENÉ STUDIUM': 'ne', 'POVINNOST JPZ': 1,
        'ZAMĚŘENÍ OBORU': 'Přírodovědné', 'ČJ+MA - % SKÓR - PRŮMĚR (PŘIJATI)': 120,
        'ČJ - % SKÓR - PRŮMĚR (PŘIJATI)': 70, 'MA - % SKÓR - PRŮMĚR (PŘIJATI)': 50,
        'KAPACITA': 10, 'PŘIHLÁŠKY CELKEM': 6, 'PŘIJATÍ': 4,
        'PŘIHLÁŠKY - PRIORITA 1': 3, 'PŘIHLÁŠKY - PRIORITA 2': 2, 'PŘIHLÁŠKY - PRIORITA 3': 1,
        'PŘIHLÁŠKY - PRIORITA 4': 0, 'PŘIHLÁŠKY - PRIORITA 5': 0,
        'NÁZEV ŠKOLY': 'Gymnázium', 'OBOR - NÁZEV': 'Gymnázium', 'TYP ŠKOLY': 'GY4',
        'OBEC': 'Praha', 'KRAJ': 'CZ010', 'KRAJ - NÁZEV': 'Praha', 'DÉLKA STUDIA': 4,
        **overrides,
    }


class TestImportCermatResults(unittest.TestCase):
    """Čtení souhrnového souboru, párování zaměření a kontroly rozsahů."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_workbook_with_explanations_and_reordered_columns(self):
        row = record()
        headers = list(reversed(row))
        wb = openpyxl.Workbook()
        wb.active.append(headers)
        wb.active.append([row[h] for h in headers])
        wb.create_sheet('vysvetlivky')
        file = self.dir / 'data.xlsx'
        wb.save(file)
        data = extract_current_year(load_flat_xlsx(file))
        result = data['123_79-41-K/41_prirodovedne']
        self.assertEqual((result['cj_ma_prijati'], result['cj_prijati'], result['ma_prijati']), (60, 35, 25))
        self.assertEqual(result['prijati'], 4)

    def test_wrong_headers_stop_import(self):
        wb = openpyxl.Workbook()
        wb.active.append(['REDIZO', 'neznámý sloupec'])
        file = self.dir / 'bad.xlsx'
        wb.save(file)
        with self.assertRaisesRegex(ValueError, 'hlavičky'):
            load_flat_xlsx(file)

    def test_duplicate_focus_stops_import(self):
        with self.assertRaisesRegex(ValueError, 'Duplicitní'):
            extract_current_year([record(), record()])

    def test_missing_score_retained_in_applications_only(self):
        row = record(**{'ČJ+MA - % SKÓR - PRŮMĚR (PŘIJATI)': None})
        self.assertEqual(extract_current_year([row]), {})
        self.assertEqual(len(build_applications([row], [], [])), 1)

    def test_priorities_not_shifted_by_results_column(self):
        result = build_applications([record()], [], [dict(id='123_79-41-K/41_Přírodovědné')])[0]
        self.assertEqual(result['pp'], [3, 2, 1, 0, 0])
        self.assertEqual(result['prihlasky'], 6)
        self.assertEqual(result['id'], '123_79-41-K/41_Přírodovědné')
        self.assertFalse(result.get('is_new'))
        with self.assertRaisesRegex(ValueError, 'priority'):
            build_applications([record(**{'PŘIHLÁŠKY CELKEM': 99})], [], [])

    def test_different_focus_is_not_matched_to_history(self):
        result = build_applications([record()], [], [dict(id='123_79-41-K/41_jazykove')])[0]
        self.assertTrue(result['is_new'])
        self.assertTrue(result['id'].endswith('_prirodovedne'))

    def test_ambiguous_history_is_not_used(self):
        old = dict(id='123_79-41-K/41_prirodovedne')
        self.assertTrue(build_applications([record()], [], [old, old])[0]['is_new'])

    def test_filter_and_range(self):
        self.assertFalse(is_valid_flat(record(**{'POVINNOST JPZ': 0})))
        self.assertFalse(is_valid_flat(record(**{'FORMA VZDĚLÁVÁNÍ': 'večerní'})))
        with self.assertRaisesRegex(ValueError, 'rozsah'):
            extract_current_year([record(**{'ČJ+MA - % SKÓR - PRŮMĚR (PŘIJATI)': 201})])

    def test_ranks_within_type(self):
        data = {key: {'school_type': typ, 'cj_ma_prijati': score}
                for key, typ, score in [('a', 'GY4', 90), ('b', 'GY4', 80), ('c', 'SOS', 70)]}
        compute_ranks(data)
        self.assertEqual(data['a']['rank_in_type'], 1)
        self.assertEqual(data['c']['rank_in_type'], 1)
        self.assertEqual(data['b']['rank_in_type'], 2)
        self.assertEqual(data['c']['type_total'], 1)

    def test_key_normalization(self):
        self.assertEqual(make_key('123', '79-41-K/41', 'IT & Sítě'), '123_79-41-K/41_it_site')

    def test_admission_context_preserves_missing_and_valid_zero(self):
        from refresh_cermat_data import admission_context
        row = record(**{'ČJ+MA - KONALI': 6, 'ČJ+MA - KONALI (PŘIJATI)': 4,
                        'ČJ+MA - % SKÓR - PRŮMĚR': 140,
                        'NEPŘIJATI - PŘIJAT NA VYŠŠÍ PRIORITU': 2,
                        'NEPŘIJATI - NEDOSTATEČNÁ KAPACITA': 0,
                        'NEPŘIJATI - NESPLNĚNÍ PODMÍNEK': 0, 'NEPŘIJATI - VZDAL SE PŘIJETÍ': 0})
        ctx = admission_context(row)
        self.assertEqual(ctx['average_all'], 70)
        self.assertEqual(ctx['average_accepted'], 60)
        self.assertEqual(ctx['capacity_rejected'], 0)
        self.assertTrue(ctx['outcomes_complete'])

        # Chybějící údaj není nula: zůstane None a rozpad se označí za neúplný.
        row['NEPŘIJATI - NEDOSTATEČNÁ KAPACITA'] = None
        self.assertIsNone(admission_context(row)['capacity_rejected'])
        self.assertFalse(admission_context(row)['outcomes_complete'])

        row['ČJ+MA - KONALI (PŘIJATI)'] = 5
        self.assertIsNone(admission_context(row)['average_accepted'])
        row['ČJ+MA - KONALI'] = 0
        self.assertIsNone(admission_context(row)['average_all'])
        row['ČJ+MA - % SKÓR - PRŮMĚR'] = 201
        with self.assertRaises(ValueError):
            admission_context(row)


if __name__ == '__main__':
    unittest.main()
